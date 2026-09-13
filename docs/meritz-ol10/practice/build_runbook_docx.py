from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\ekzm8\OneDrive\바탕 화면\OneDrive\문서\등촌프로젝트\docs\meritz-ol10\practice")
SRC = ROOT / "MERITZ_OL10_FIELD_RUNBOOK_DRAFT.md"
OUT = ROOT / "MERITZ_OL10_FIELD_RUNBOOK.docx"

ENV = {
    "2. 시작 점검": "APP IP, DB IP, SSH 포트, DB 실제 listen 포트를 먼저 기록한다. 여기 값이 이후 모든 연결 시험의 기준이다.",
    "3. MariaDB - DB 서버": "DB IP와 DB port, DB명, 애플리케이션 접속 계정, APP 대역 방화벽 허용값을 대상 환경에 맞춘다.",
    "4. APP runtime과 권한 - APP 서버": "APP 서버의 실제 설치 경로, 실행 계정, admin UID/GID를 확인한다. 원본과 다르면 권한을 맞춘 뒤 진행한다.",
    "5. Redis": "Redis bind 주소, cluster-announce-ip, 7000/7001/7002 포트와 데이터 경로를 대상 APP IP로 맞춘다. RDB 복원 뒤 AOF 정책을 확인한다.",
    "6. ZooKeeper": "dataDir/dataLogDir, clientPort 2181/2182/2183, server.1~3의 대상 APP IP와 myid 1/2/3을 맞춘다. 기동 전에 로그 디렉터리를 만든다.",
    "7. SolrCloud": "SOLR_HOST, SOLR_PORT 8983/8984/8985, SOLR_LOGS_DIR, SOLR_DATA_HOME, ZK_HOST의 IP와 마지막 /solr를 확인한다. /solr와 configset은 컬렉션보다 먼저 준비한다.",
    "8. Elasticsearch와 Logstash": "ES node.name은 유지하고 network.publish_host만 대상 APP IP로 변경한다. http.port 9200/9201/9202, path.data, path.repo는 실제 경로와 일치해야 한다. Logstash 입력 파일과 pipeline 경로도 확인한다.",
    "9. YML, build, 애플리케이션": "DB JDBC URL과 port, Redis nodes, Solr URL, ZK_HOST, ES hosts, engine/master/scheduler/chat-ui 내부 URL을 대상 환경에 맞춘다. 소스 YML과 JAR 내부 설정을 모두 확인한다.",
    "10. 최종 시험": "브라우저용 hosts 이름과 실제 접속 주소, CMS/Chat UI 포트, 테스트 botCode와 alias를 대상 환경에 맞춘다.",
}

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def borders(table, color='D9D9D9'):
    tbl = table._tbl
    tblPr = tbl.tblPr
    b = tblPr.first_child_found_in('w:tblBorders')
    if b is None:
        b = OxmlElement('w:tblBorders')
        tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = b.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            b.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)

def set_font(run, name='Malgun Gothic', size=10.2, bold=False, color='000000'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_text(doc, text, style=None, size=10.2, bold=False, color='000000', space_after=3):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return p

def add_code(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.05)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        set_font(r, name='Consolas', size=8.4, color='1F1F1F')

def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders(table)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    trPr.append(header)
    for j, val in enumerate(rows[0]):
        c = table.rows[0].cells[j]
        c.text = str(val).strip()
        shade(c, '1F4E78')
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            for r in p.runs: set_font(r, size=8.5, bold=True, color='FFFFFF')
    for i, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = str(val).strip()
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2 == 1: shade(cells[j], 'F2F6FA')
            for p in cells[j].paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs: set_font(r, size=8.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def main():
    text = SRC.read_text(encoding='utf-8')
    raw = text.splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.58); sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.65); sec.right_margin = Inches(0.65)
    styles = doc.styles
    for s in ('Normal','Body Text'):
        styles[s].font.name = 'Malgun Gothic'; styles[s].font.size = Pt(10.2)
        styles[s]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    for s in ('Title','Heading 1','Heading 2','Heading 3'):
        styles[s].font.name = 'Malgun Gothic'; styles[s].font.color.rgb = RGBColor(0,0,0)
        styles[s]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    styles['Title'].font.size = Pt(23)
    styles['Heading 1'].font.size = Pt(15); styles['Heading 1'].font.bold = True
    styles['Heading 2'].font.size = Pt(11.5); styles['Heading 2'].font.bold = True

    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run('Meritz OL10 현장 이관 실행서')
    set_font(r, size=23, bold=True)
    p.paragraph_format.space_after = Pt(2)
    add_text(doc, '2회차 실전 운영 체크리스트', size=12, bold=True, color='1F4E78', space_after=10)
    add_text(doc, '목적: 이미 동작하는 APP와 DB 환경을 다른 서버로 옮길 때, 환경값 확인과 서비스 검증을 빠뜨리지 않고 6시간 안에 순서대로 진행하기 위한 현장용 문서다.', size=10.5, space_after=4)
    add_text(doc, '사용 원칙: 명령은 표시된 서버에서 실행한다. APP와 DB 주소, 포트, 경로는 고정값으로 외우지 말고 현장값 표에 먼저 적은 뒤 각 절차의 환경값 문장을 기준으로 바꾼다. 비밀번호와 개인키는 이 문서에 기록하지 않는다.', size=10.5, space_after=8)

    # parse markdown
    i=0; in_code=False; code=[]; table=[]; current=None
    while i < len(raw):
        line = raw[i]
        if line.startswith('```'):
            if in_code:
                add_code(doc, code); code=[]; in_code=False
            else: in_code=True
            i+=1; continue
        if in_code:
            code.append(line); i+=1; continue
        if line.startswith('# ') and not line.startswith('# 부록'):
            i+=1; continue
        if line.startswith('## '):
            if table: add_table(doc, table); table=[]
            heading=line[3:].strip(); current=heading
            p=doc.add_paragraph(style='Heading 1'); p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(4)
            r=p.add_run(heading); set_font(r,size=15,bold=True)
            if heading in ENV:
                p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4)
                r=p.add_run('이 단계에서 반드시 확인하거나 바꿀 환경값: '); set_font(r,size=9.6,bold=True,color='1F4E78')
                r=p.add_run(ENV[heading]); set_font(r,size=9.6,color='1F4E78')
            i+=1; continue
        if line.startswith('### '):
            if table: add_table(doc, table); table=[]
            p=doc.add_paragraph(style='Heading 2'); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(3)
            r=p.add_run(line[4:].strip()); set_font(r,size=11.5,bold=True)
            i+=1; continue
        if line.startswith('|'):
            table.append([x.strip() for x in line.strip('|').split('|')])
            i+=1
            if i < len(raw) and re.match(r'^\s*\|?\s*:?-+:?', raw[i]): i+=1
            continue
        if table:
            add_table(doc, table); table=[]
        if not line.strip():
            i+=1; continue
        if line.startswith('    '):
            # indented non-fenced command block
            code_lines=[]
            while i < len(raw) and (raw[i].startswith('    ') or not raw[i].strip()):
                if raw[i].startswith('    '): code_lines.append(raw[i][4:])
                i+=1
            add_code(doc, code_lines); continue
        if line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2)
            r=p.add_run(line[2:].strip()); set_font(r,size=9.8)
            i+=1; continue
        if re.match(r'^\[[ xX]\]', line.strip()):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.12); p.paragraph_format.space_after=Pt(2)
            r=p.add_run(line.strip()); set_font(r,size=9.8)
            i+=1; continue
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.05
        r=p.add_run(line.strip()); set_font(r,size=9.8)
        i+=1
    if table: add_table(doc, table)
    if code: add_code(doc, code)

    # footer
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('Meritz OL10 현장 이관 실행서 | 실전용')
    set_font(r,size=8,color='666666')
    doc.core_properties.title = 'Meritz OL10 현장 이관 실행서'
    doc.core_properties.subject = '2회차 실전 운영 체크리스트'
    doc.core_properties.author = 'Codex'
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__': main()
