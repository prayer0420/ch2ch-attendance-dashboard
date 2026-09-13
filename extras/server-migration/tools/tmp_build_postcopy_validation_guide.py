from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "산출물"
OUT.mkdir(parents=True, exist_ok=True)
FINAL = OUT / "메리츠_운영_이관완료후_검증_및_연결설정_가이드.docx"


def set_font(run, name="맑은 고딕", size=None, bold=None, color=(0, 0, 0)):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rpr.rFonts.set(qn(key), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(*color)


def style_font(style, size, bold=False):
    style.font.name = "맑은 고딕"
    rpr = style._element.get_or_add_rPr()
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rpr.rFonts.set(qn(key), "맑은 고딕")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def border(cell, color="D9D9D9"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = borders.find(qn("w:" + edge))
        if item is None:
            item = OxmlElement("w:" + edge)
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "6")
        item.set(qn("w:color"), color)


def margins(cell, value=100):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for side in ("top", "start", "bottom", "end"):
        node = mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


class Guide:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11)
        sec.top_margin = Cm(1.65)
        sec.bottom_margin = Cm(1.55)
        sec.left_margin = Cm(1.75)
        sec.right_margin = Cm(1.75)
        sec.footer_distance = Cm(0.7)

        styles = self.doc.styles
        style_font(styles["Normal"], 9.5)
        styles["Normal"].paragraph_format.line_spacing = 1.13
        styles["Normal"].paragraph_format.space_after = Pt(4)
        style_font(styles["Title"], 23, True)
        styles["Title"].paragraph_format.space_after = Pt(10)
        ppr = styles["Title"]._element.get_or_add_pPr()
        pbd = ppr.find(qn("w:pBdr"))
        if pbd is not None:
            ppr.remove(pbd)
        for name, size, before, after in [
            ("Heading 1", 16, 15, 7),
            ("Heading 2", 12.5, 10, 5),
            ("Heading 3", 10.5, 7, 3),
        ]:
            style_font(styles[name], size, True)
            styles[name].paragraph_format.space_before = Pt(before)
            styles[name].paragraph_format.space_after = Pt(after)
            styles[name].paragraph_format.keep_with_next = True

        if "Code Block" not in styles:
            code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        else:
            code = styles["Code Block"]
        code.font.name = "Consolas"
        code._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
        code.font.size = Pt(7.7)
        code.paragraph_format.left_indent = Cm(0.25)
        code.paragraph_format.right_indent = Cm(0.15)
        code.paragraph_format.space_before = Pt(2)
        code.paragraph_format.space_after = Pt(5)
        code.paragraph_format.line_spacing = 1.0

        title = self.doc.add_paragraph(style="Title")
        title.add_run("메리츠 운영 이관 완료 후 검증 및 연결 설정 가이드")
        ppr = title._p.get_or_add_pPr()
        pbd = ppr.find(qn("w:pBdr"))
        if pbd is not None:
            ppr.remove(pbd)
        sub = self.doc.add_paragraph()
        set_font(sub.add_run("프로그램과 데이터 복사가 완료된 운영 서버에서 연결 설정과 정상 동작을 확인하는 절차"), size=11, bold=True)
        self.p("적용 기준 2026년 9월 10일 1차 이관 실습 결과")
        self.p("보안 원칙 비밀번호 인증 키 토큰은 화면 또는 문서에 기록하지 않는다")
        self.p("중요 이 문서는 재복사나 클러스터 재생성을 기본 절차로 다루지 않는다. 운영 환경의 실제 IP 계정 포트와 서비스 배치가 승인된 구성표와 일치하는지 먼저 확인한다.")

        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(footer.add_run("메리츠 운영 이관 후 검증 및 연결 설정"), size=8, color=(110, 110, 110))

    def h1(self, text):
        return self.doc.add_heading(text, level=1)

    def h2(self, text):
        return self.doc.add_heading(text, level=2)

    def h3(self, text):
        return self.doc.add_heading(text, level=3)

    def p(self, text, lead=None):
        p = self.doc.add_paragraph()
        if lead and text.startswith(lead):
            r = p.add_run(lead)
            set_font(r, bold=True)
            p.add_run(text[len(lead):])
        else:
            p.add_run(text)
        return p

    def bullets(self, items):
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.add_run(item)

    def code(self, text):
        p = self.doc.add_paragraph(style="Code Block")
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F3F5F7")
        ppr.append(shd)
        lines = text.strip("\n").splitlines()
        for i, line in enumerate(lines):
            r = p.add_run(line)
            set_font(r, "Consolas", 7.35 if len(line) > 110 else 7.7)
            if i < len(lines) - 1:
                r.add_break()
        return p

    def table(self, headers, rows, widths, font=8.2, center_cols=(0,)):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Table Grid"
        for i, text in enumerate(headers):
            c = table.rows[0].cells[i]
            c.width = Cm(widths[i])
            shade(c, "1F4E78")
            border(c)
            margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(str(text)), size=8.4, bold=True, color=(255, 255, 255))
        trpr = table.rows[0]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:tblHeader"))
        trpr.append(OxmlElement("w:cantSplit"))
        for ridx, row in enumerate(rows):
            added_row = table.add_row()
            added_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            cells = added_row.cells
            for i, text in enumerate(row):
                c = cells[i]
                c.width = Cm(widths[i])
                if ridx % 2:
                    shade(c, "F6F9FC")
                border(c)
                margins(c)
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT
                set_font(p.add_run(str(text)), size=font)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)
        return table

    def save(self):
        self.doc.save(FINAL)


g = Guide()

g.h1("1 문서 범위와 작업 순서")
g.p("운영 이관에서는 프로그램 디렉터리와 데이터가 대상 서버에 이미 복사되어 있다고 가정한다. 따라서 작업자는 다시 복사하기보다 복사 결과를 확인하고, 운영 주소에 맞게 노드별 설정을 수정한 뒤, 하위 인프라부터 상위 애플리케이션 순서로 연결을 검증한다.")
g.table(["단계", "작업", "완료 기준"], [
    ["1", "서버 식별과 파일 확인", "hostname IP 디렉터리 소유권이 승인표와 일치"],
    ["2", "이름 해석과 방화벽", "서비스 이름이 운영 IP를 반환하고 필수 포트 통신 가능"],
    ["3", "데이터 계층 확인", "MariaDB Redis ZooKeeper Solr Elasticsearch 정상"],
    ["4", "애플리케이션 설정", "DB Redis Solr Elasticsearch 내부 API 주소가 운영값"],
    ["5", "순차 기동", "프로세스와 LISTEN 포트 정상"],
    ["6", "웹과 업무 기능 검증", "로그인 조회 대화 학습 배포가 승인 기준 충족"],
], [1.4, 6.0, 8.0])
g.p("금지 사항: 확인 없이 Redis cluster create, FLUSHDB, MariaDB 계정 DDL, Solr RESTORE, Elasticsearch restore를 다시 실행하지 않는다. 이미 복사된 운영 데이터가 있으면 중복 생성 또는 덮어쓰기가 발생할 수 있다.")

g.h2("운영값 기록표")
g.table(["구분", "노드 1", "노드 2", "노드 3", "비고"], [
    ["호스트명", "", "", "", "승인된 FQDN"],
    ["운영 IP", "", "", "", "실습값 81 82 83으로 대체 금지"],
    ["MariaDB 역할", "Master", "Replica", "Replica", "실제 역할 확인"],
    ["Redis 별칭", "redis1", "redis2", "redis3", ""],
    ["ZooKeeper myid", "1", "2", "3", ""],
    ["Solr 별칭", "solr1", "solr2", "solr3", ""],
    ["Elasticsearch 별칭", "es1", "es2", "es3", "메리츠 7.8.0"],
], [3.2, 3.0, 3.0, 3.0, 4.1])

g.h1("2 공통 사전 확인과 연결 기반")
g.h2("서버와 복사 결과 확인")
g.code("hostname\nhostname -i\ndate '+%F %T %Z'\ndf -hT / /data /application /tmp 2>/dev/null\nls -ld /application /logs /data 2>/dev/null")
g.p("정상 기준: 접속 서버가 작업 대상과 일치하고, 프로그램과 데이터 경로에 충분한 여유 공간이 있으며, 실행 계정이 필요한 디렉터리를 읽고 쓸 수 있다.")
g.code("for d in mariadb redis solr elk gateway engine master cms chat-ui scheduler; do\n  [ -e \"/application/$d\" ] && ls -ld \"/application/$d\" || echo \"$d MISSING\"\ndone")
g.p("MISSING이 반드시 오류는 아니다. 승인된 서비스 배치표에 해당 서버에 있어야 하는 서비스인지 대조한다.")

g.h2("이름 해석")
g.p("호스트명은 두 종류를 함께 사용한다. 번호가 붙은 이름은 실제 노드, 번호가 없는 이름은 애플리케이션이 사용하는 논리 주소다. 번호 이름을 지우거나 번호 없는 이름만 세 노드에 중복 등록하면 장애 원인 추적과 장애조치가 어려워진다.")
g.table(["구분", "예", "용도", "가리킬 대상"], [
    ["논리 이름", "engine gateway cms master chat-ui", "애플리케이션 설정과 브라우저 진입", "승인된 VIP 또는 HAProxy"],
    ["노드 이름", "engine1~3 gateway1~3 cms1~3", "HAProxy backend와 노드 직접 점검", "각 물리 노드 IP"],
    ["인프라 이름", "mariadb redis1~3 zk1~3 solr1~3 es1~3", "데이터 계층 연결", "승인된 서비스 IP"],
], [3.0, 4.6, 5.1, 4.6], font=7.7)
g.code("getent hosts engine gateway cms master chat-ui\ngetent hosts engine1 engine2 engine3 gateway1 gateway2 gateway3\ngetent hosts cms1 cms2 cms3 master1 master2 master3 chat-ui1 chat-ui2 chat-ui3\ngetent hosts mariadb redis1 redis2 redis3 zk1 zk2 zk3 solr1 solr2 solr3 es1 es2 es3")
g.p("정상 기준: 모든 이름이 승인된 운영 IP를 반환한다. 실습에서는 무번호 이름을 192.168.1.81의 HAProxy로 보냈지만, 운영에서는 승인된 VIP 또는 외부 로드밸런서를 사용한다. source xen02와 target 운영계가 동시에 켜져 있을 때 동일한 별칭을 양쪽 IP에 중복 등록하지 않는다.")
g.p("설정 위치: 서버는 /etc/hosts 또는 운영 DNS를 사용한다. 작업자 PC도 CMS와 시뮬레이터가 실제로 참조하는 cms, chat-ui, gateway를 해석할 수 있어야 한다. Windows hosts 변경 후 ipconfig /flushdns를 실행한다.")
g.code("# 실습 예시이며 운영에서는 승인된 VIP와 실제 노드 IP로 교체\n192.168.1.81 engine gateway cms master chat-ui\n192.168.1.81 engine1 gateway1 cms1 master1 chat-ui1\n192.168.1.82 engine2 gateway2 cms2 master2 chat-ui2\n192.168.1.83 engine3 gateway3 cms3 master3 chat-ui3\n\n# Windows 작업자 PC 예시\n192.168.1.81 cms chat-ui gateway engine master")

g.h2("포트와 방화벽")
g.table(["서비스", "포트", "확인 대상"], [
    ["HAProxy", "80", "작업자 PC와 내부 사용자"],
    ["MaxScale", "3306", "애플리케이션"],
    ["MariaDB", "13306", "MaxScale과 복제 노드"],
    ["Redis", "7000 7001 17000 17001", "3개 Redis 노드 상호 간"],
    ["ZooKeeper", "2181 2888 3888", "3개 ZK 노드와 Solr"],
    ["SolrCloud", "8983", "애플리케이션과 Solr 노드 간"],
    ["메리츠 Elasticsearch", "9210 9310", "클라이언트와 ES 노드 간"],
    ["Applications", "8081 8180 8280 8380 8480 8580", "HAProxy와 내부 API"],
], [4.0, 4.2, 8.1])
g.code("sudo firewall-cmd --get-active-zones\nsudo firewall-cmd --zone=public --list-ports\nsudo firewall-cmd --zone=public --list-rich-rules\nsudo ss -lntp\n\n# 노드 간 Engine 통신 예시\nfor h in engine1 engine2 engine3; do nc -zvw3 \"$h\" 8180; done")
g.p("애플리케이션 프로세스가 떠 있어도 노드 간 8180이 막혀 있으면 CMS 학습 배포가 일부 Engine에 전달되지 않는다. 운영 보안정책에 따라 HAProxy/VIP와 애플리케이션 노드 사이의 8081, 8180, 8280, 8380, 8480, 8580을 필요한 방향으로 허용한다. 전체 공개보다 승인된 애플리케이션 서브넷 또는 노드 IP를 source로 제한한다.")
g.code("# 예: 승인된 애플리케이션 대역에서 Engine 8180 허용\nsudo firewall-cmd --permanent --zone=public \\\n  --add-rich-rule='rule family=\"ipv4\" source address=\"<APP_SUBNET>\" port port=\"8180\" protocol=\"tcp\" accept'\nsudo firewall-cmd --reload\nfor h in engine1 engine2 engine3; do nc -zvw3 \"$h\" 8180; done")
g.p("No route to host 또는 TIMEOUT은 방화벽이나 경로 문제일 가능성이 높고, Connection refused는 주소에는 도달했지만 서비스가 LISTEN하지 않는 상태다. ping 성공만으로 TCP와 HTTP 연결 성공을 판단하지 않는다.")

g.h1("3 MariaDB와 MaxScale")
g.h2("확인할 것")
g.bullets([
    "세 노드의 server id가 서로 다르고 역할이 승인된 구성과 일치하는지 확인한다.",
    "Replica 두 대의 IO와 SQL thread, 지연, 마지막 SQL 오류를 확인한다.",
    "MaxScale 서비스 계정은 사용자명과 접속 Host 조합이 실제 MaxScale IP와 일치해야 한다.",
    "애플리케이션은 개별 MariaDB가 아니라 MaxScale 별칭 mariadb와 3306을 사용한다.",
])
g.h2("노드와 복제 검증")
g.code("/mariadb/bin/mariadb -S /mariadb/run/mariadb.sock -e \\\n\"SELECT @@hostname,@@server_id,@@read_only; SHOW SLAVE STATUS\\G\" | \\\ngrep -E 'hostname|server_id|read_only|Master_Host|Slave_IO_Running|Slave_SQL_Running|Seconds_Behind_Master|Last_SQL_Error|Gtid_IO_Pos'")
g.p("정상 기준: Replica에서 Slave IO Running과 Slave SQL Running이 Yes, Seconds Behind Master가 0, Last SQL Error가 빈 값이다. Master에는 SHOW SLAVE STATUS 결과가 없을 수 있다.")
g.h2("계정과 MaxScale 연결")
g.code("/mariadb/bin/mariadb -S /mariadb/run/mariadb.sock -e \\\n\"SELECT User,Host,plugin FROM mysql.user WHERE User='mxs_service';\"\n\nsudo maxctrl list servers\nnc -zvw3 mariadb 3306")
g.p("설정할 것: 계정 Host는 MaxScale이 DB에 접속할 때 보이는 실제 IP로 생성한다. MaxScale 서비스의 user와 password는 DB 계정과 동일하게 설정한다. 비밀번호는 명령 이력에 직접 남기지 않는 방식을 사용한다.")
g.p("정상 기준: MaxScale에서 1대는 Master Running, 2대는 Slave Running이고 애플리케이션 서버에서 mariadb 3306 연결이 된다.")
g.p("중지 조건: Last SQL Error, GTID out of order, Slave SQL Running No가 있으면 애플리케이션 기동 전에 복제를 먼저 복구한다.")

g.h1("4 Redis Cluster")
g.h2("설정할 것")
g.table(["항목", "노드별 값", "주의점"], [
    ["cluster announce ip", "redis1 redis2 redis3", "각 서버 자신을 가리켜야 함"],
    ["client ports", "7000 7001", "두 프로세스 모두 실행"],
    ["bus ports", "17000 17001", "노드 상호 통신 허용"],
    ["data dirs", "/application/redis/db/7000 및 7001", "admin 쓰기 권한"],
    ["log dir", "/logs/redis", "admin 쓰기 권한"],
], [5.0, 6.0, 5.3])
g.code("grep -E '^(port|dir|cluster-enabled|cluster-config-file|cluster-announce-ip|cluster-announce-port|cluster-announce-bus-port)' \\\n/application/redis/conf/7000.conf /application/redis/conf/7001.conf\n\n/application/redis/status.sh")
g.h2("클러스터 연결 검증")
g.code("for h in redis1 redis2 redis3; do\n  for p in 7000 7001 17000 17001; do nc -zvw2 \"$h\" \"$p\"; done\ndone\n\n/application/redis/bin/redis-cli --cluster check redis1:7000\n/application/redis/bin/redis-cli -h redis1 -p 7000 cluster info\n/application/redis/bin/redis-cli -h redis1 -p 7000 cluster nodes")
g.p("정상 기준: cluster state ok, known nodes 6, cluster size 3, 16384 slots covered, master 3대와 replica 3대가 connected 상태다.")
g.p("운영 판단: 복사된 nodes.conf와 데이터가 유효하면 cluster create를 다시 실행하지 않는다. 노드 주소가 source IP를 계속 광고하면 복사 방식과 재구성 계획을 별도 승인받는다. Redis cache는 재생성 정책이면 key 수 불일치만으로 장애로 판단하지 않는다.")

g.h1("5 ZooKeeper")
g.h2("설정할 것")
g.code("grep -Ev '^[[:space:]]*(#|$)' /application/solr/zk/conf/zoo.cfg\ncat /application/solr/zk/data/myid\nls -ld /application/solr/zk/data /application/solr/zk/datalog /logs/solr/zk")
g.table(["노드", "myid", "server 항목"], [
    ["zk1", "1", "server.1=zk1:2888:3888"],
    ["zk2", "2", "server.2=zk2:2888:3888"],
    ["zk3", "3", "server.3=zk3:2888:3888"],
], [4.2, 3.0, 9.1])
g.p("dataDir, dataLogDir, 로그 디렉터리는 실행 계정 admin이 쓸 수 있어야 한다. myid는 각 서버마다 달라야 하고 zoo.cfg의 server 번호와 일치해야 한다.")
g.h2("연결과 데이터 검증")
g.code("for h in zk1 zk2 zk3; do\n  echo \"===== $h =====\"\n  printf 'ruok' | nc -w 3 \"$h\" 2181; echo\n  printf 'srvr' | nc -w 3 \"$h\" 2181 | grep -E 'Mode|Node count'\ndone")
g.p("nc가 없으면 bash의 /dev/tcp를 사용한다.")
g.code("timeout 3 bash -c 'exec 3<>/dev/tcp/zk1/2181; printf \"srvr\" >&3; cat <&3' | grep -E 'Mode|Node count'")
g.p("정상 기준: 세 노드 모두 imok, leader 1대, follower 2대, Node count가 동일하다. 데이터가 복사된 운영 환경에서는 source와 target의 znode 또는 Solr chroot 구성이 승인된 결과와 일치하는지도 확인한다.")

g.doc.add_page_break()
g.h1("6 SolrCloud")
g.h2("노드별 연결 설정")
g.code("grep -nE '^(SOLR_HOST|SOLR_PORT|ZK_HOST)' /application/solr/solr/bin/solr.in.sh")
g.table(["노드", "SOLR HOST", "SOLR PORT", "ZK HOST"], [
    ["노드 1", "solr1", "8983", "zk1:2181,zk2:2181,zk3:2181/운영 chroot"],
    ["노드 2", "solr2", "8983", "동일"],
    ["노드 3", "solr3", "8983", "동일"],
], [3.0, 3.5, 3.0, 6.8])
g.p("설정할 것: 세 노드의 SOLR HOST는 각각 달라야 하며 ZK HOST는 동일해야 한다. source와 target이 같은 ZooKeeper를 공유하는 구조라면 target 전용 chroot를 사용한다. 실습에서는 solr-xen01을 사용했다.")
g.h2("프로세스와 클러스터 검증")
g.code("for ip in <운영_IP_1> <운영_IP_2> <운영_IP_3>; do\n  curl -sS --max-time 10 \"http://$ip:8983/solr/admin/info/system?wt=json\" | \\\n  grep -E '\"mode\"|\"zkHost\"|\"solr-spec-version\"'\ndone\n\ncurl -sS 'http://solr1:8983/solr/admin/collections?action=LIST&wt=json'\ncurl -sS 'http://solr1:8983/solr/admin/collections?action=LISTALIASES&wt=json'\ncurl -sS 'http://solr1:8983/solr/admin/collections?action=CLUSTERSTATUS&wt=json' | \\\ngrep -oE '\"state\":\"[^\"]+\"' | sort | uniq -c")
g.p("정상 기준: 세 노드 모두 mode solrcloud, 동일한 운영 ZK chroot, live nodes 3대, 필요한 collection과 alias 존재, replica가 모두 active다.")
g.h2("복사 데이터 확인")
g.code("for c in TC_CALLBOT_LOG TC_CALLBOT_B tc_01_A TC_CALLBOT_TEST tc_01_B TC_CALLBOT_A; do\n  printf '%-24s ' \"$c\"\n  curl -sS --max-time 30 \"http://solr1:8983/solr/$c/select?q=*:*&rows=0&wt=json\" | \\\n  grep -oE '\"numFound\":[0-9]+'\ndone")
g.p("정상 기준: 승인된 source 기준 문서 수와 일치한다. alias TC CALLBOT과 tc 01이 각각 올바른 A 또는 B collection을 가리키는지도 확인한다. 복사된 데이터가 정상인데 collection을 다시 RESTORE하지 않는다.")

g.h1("7 Elasticsearch")
g.h2("운영 구조")
g.p("메리츠용 Elasticsearch 7.8.0과 기존 Elasticsearch 9.1.1을 같은 서버에서 동시에 유지하는 경우에만 포트 데이터 로그 경로를 완전히 분리한다. 1차 실습에서는 충돌 회피를 위해 메리츠 7.8.0을 9210/9310으로 변경했다. 실제 개발계에 충돌이 없다면 승인된 원래 포트 9200/9300을 사용하고 애플리케이션 설정도 그 값에 맞춘다.")
g.table(["구분", "HTTP", "Transport", "data", "logs"], [
    ["기존 9.1.1", "9200", "9300", "/var/lib/elasticsearch", "/var/log/elasticsearch"],
    ["메리츠 7.8.0", "9210", "9310", "/application/elk-meritz/data", "/logs/elk-meritz"],
], [3.2, 2.3, 2.8, 4.4, 4.5])
g.h2("노드별 설정")
g.code("grep -nE '^(cluster.name|node.name|path.data|path.logs|path.repo|network.host|http.port|transport.port|discovery.seed_hosts)' \\\n/application/elk/config/elasticsearch.yml\n\ngrep -nE 'HeapDumpPath|ErrorFile|Xloggc|Xlog:gc' /application/elk/config/jvm.options")
g.p("설정할 것: cluster name은 trusted-context-es, node name은 es1 es2 es3, network host는 각 운영 IP, discovery seed hosts는 es1:9310부터 es3:9310까지다. JVM 로그 경로도 /logs/elk-meritz로 분리한다. path.repo는 snapshot을 사용할 때만 모든 노드에 동일하게 적용한다.")
g.h2("동시 운영과 클러스터 검증")
g.code("sudo ss -lntp | grep -E ':(9200|9300|9210|9310)\\b'\n\ncurl -sS 'http://es1:9210/_cat/nodes?v&h=name,ip,role,master,version'\ncurl -sS 'http://es1:9210/_cluster/health?pretty'\ncurl -sS 'http://es1:9210/_cat/indices?expand_wildcards=all&v&h=health,status,index,docs.count,store.size&s=index'\n\n# 기존 9.1.1을 유지하는 노드에서 별도 확인\ncurl -sS 'http://127.0.0.1:9200/' | grep -E '\"name\"|\"cluster_name\"|\"number\"'")
g.p("정상 기준: 메리츠 클러스터는 es1부터 es3까지 3 nodes, version 7.8.0, health green, 인덱스와 문서 수가 source 기준과 일치한다. 기존 9.1.1도 원래 클러스터에 연결되어 있어야 한다.")
g.p("주의: network host가 서버 IP로 지정되면 127.0.0.1 요청은 거부될 수 있다. 이 경우 실제 서버 IP 또는 es1 별칭으로 확인한다.")

g.h1("8 애플리케이션 연결 설정")
g.h2("배치 원칙")
g.p("gateway engine master cms chat-ui가 여러 노드에 복사되어 있어도 실제 운영 기동 수와 HAProxy backend는 승인된 배치표를 따른다. scheduler는 중복 작업 방지를 위해 승인된 단일 노드에서만 기동하는 것을 기본으로 한다.")
g.h2("반드시 확인하고 수정할 연결값")
g.table(["의존 서비스", "권장 연결값", "주요 확인 위치", "정상 기준"], [
    ["MariaDB", "mariadb:3306", "runtime.env application yml", "MaxScale 경유 접속"],
    ["Redis", "redis1:7000 redis2:7000 redis3:7000", "application yml", "cluster state ok"],
    ["ZooKeeper", "zk1:2181 zk2:2181 zk3:2181", "Solr 설정", "3대 quorum"],
    ["Solr", "solr1:8983 solr2:8983 solr3:8983", "application yml 또는 내부 설정", "운영 chroot의 collection"],
    ["Elasticsearch", "es1 es2 es3 port 9210", "gateway master cms scheduler", "trusted-context-es 7.8.0"],
    ["Master API", "http://master", "master cms 설정", "DNS와 8380 backend 정상"],
    ["Engine", "논리 주소 http://engine, backend engine1~3:8180", "gateway master cms", "세 backend 모두 통신 가능"],
    ["Scheduler", "scheduler1:8580/scheduler", "batch URL", "한 노드만 기동"],
], [3.2, 5.5, 4.2, 4.2], font=7.9)
g.code("grep -RniE 'mariadb|3306|redis[123]|7000|zk[123]|2181|solr[123]|8983|elasticsearch|es[123]|9200|9210|master|engine[123]|scheduler1' \\\n/application/{gateway,engine,master,cms,chat-ui,scheduler}/conf 2>/dev/null | \\\ngrep -vE 'before-|backup/'")
g.p("설정할 것: 각 연결값은 실제 운영 배치표를 기준으로 한다. 9210은 기존 9.1.1과 충돌했던 실습 전용 값이며, 운영에서 충돌이 없으면 메리츠 7.8.0의 승인 포트인 9200을 유지한다. 기존 9.1.1을 사용하는 별도 애플리케이션 설정은 변경하지 않는다. DB 계정과 비밀번호는 source의 승인된 운영값을 안전한 방식으로 반영한다.")

g.h2("로그 디렉터리와 기동 전 연결 확인")
g.code("sudo install -d -o admin -g admin -m 750 \\\n/logs/master /logs/cms /logs/engine /logs/gateway /logs/chat-ui /logs/scheduler\n\nnc -zvw2 mariadb 3306\n/application/redis/bin/redis-cli -h redis1 -p 7000 cluster info | grep cluster_state\nprintf 'ruok' | nc -w 3 zk1 2181; echo\ncurl -sS --max-time 10 'http://solr1:8983/solr/admin/info/system?wt=json' | grep '\"mode\"'\ncurl -sS --max-time 10 'http://es1:<ES_HTTP_PORT>/' | grep -E '\"cluster_name\"|\"number\"'")
g.p("다섯 의존 서비스가 모두 정상일 때 애플리케이션을 기동한다. 연결 실패 상태에서 애플리케이션부터 반복 재시작하면 원인 로그가 누적되어 판단이 어려워진다.")

g.h1("9 애플리케이션과 HAProxy 기동")
g.h2("기동 순서")
g.table(["순서", "서비스", "확인 포트", "설명"], [
    ["1", "Master", "8380", "내부 관리 API 선행"],
    ["2", "CMS", "8280", "Master 연결과 DB 초기화 확인"],
    ["3", "Engine", "8180", "대화 처리 노드"],
    ["4", "Gateway", "8081", "외부 요청 진입"],
    ["5", "Chat UI", "8480", "사용자 화면"],
    ["6", "Scheduler", "8580", "승인된 단일 노드만"],
    ["7", "HAProxy", "80", "backend 확인 후 전환"],
], [1.5, 3.5, 3.0, 8.2])
g.code("/application/master/start.sh\n/application/cms/start.sh\n/application/engine/start.sh\n/application/gateway/start.sh\n/application/chat-ui/start.sh\n/application/scheduler/start.sh")
g.h3("기동 후 상태 확인")
g.code("for d in master cms engine gateway chat-ui scheduler; do\n  echo \"===== $d =====\"\n  /application/$d/status.sh 2>&1\ndone\n\nsudo ss -lntp | grep -E ':(8081|8180|8280|8380|8480|8580)\\b'")
g.p("정상 기준: status가 running이고 각 프로세스의 LISTEN 포트가 확인된다. start 출력만 보고 성공으로 판정하지 않는다. 몇 초 후 stopped가 되면 해당 서비스의 최신 로그를 확인한다.")

g.h2("HAProxy와 웹 주소")
g.code("sudo haproxy -c -f /etc/haproxy/haproxy.cfg\nsudo systemctl restart haproxy\nsudo systemctl --no-pager --full status haproxy\nsudo ss -lntp | grep ':80\\b'\n\ncurl -sS -o /tmp/cms_login.html -w 'HTTP_CODE=%{http_code} SIZE=%{size_download}\\n' \\\n-H 'Host: cms' http://127.0.0.1/login")
g.p("정상 기준: 설정 파일 valid, HAProxy active, 80 LISTEN, CMS login HTTP 200이다. HTTP 503은 HAProxy가 살아 있지만 사용 가능한 backend를 찾지 못했다는 뜻이다.")
g.p("작업자 PC 설정: cms와 meritz.easycms뿐 아니라 시뮬레이터가 여는 chat-ui와 API 호출 대상 gateway도 운영 HAProxy 또는 VIP IP로 해석되어야 한다. DNS cache를 비운 뒤 curl 또는 브라우저로 TCP 80과 HTTP 응답을 확인한다. ping 성공은 ICMP 확인일 뿐이다.")
g.code("# Windows 관리자 CMD 또는 PowerShell\nipconfig /flushdns\nping cms\nping chat-ui\nping gateway\ncurl.exe -v --max-time 10 http://gateway/gateway/")

g.doc.add_page_break()
g.h1("10 서비스별 최종 판정표")
g.table(["확인", "서비스", "필수 판정 기준", "결과와 시간"], [
    ["[ ]", "MariaDB", "복제 IO SQL Yes 지연 0 오류 없음", ""],
    ["[ ]", "MaxScale", "1 Master 2 Slave Running mariadb 3306 연결", ""],
    ["[ ]", "Redis", "6 nodes 3 masters 3 replicas 16384 slots", ""],
    ["[ ]", "ZooKeeper", "3 imok leader 1 follower 2 Node count 동일", ""],
    ["[ ]", "Solr", "3 live nodes collection alias replica 문서 수 정상", ""],
    ["[ ]", "ES 7.8", "3 nodes green indices와 문서 수 일치", ""],
    ["[ ]", "ES 9.1", "기존 9200 9300 서비스 유지", ""],
    ["[ ]", "Applications", "승인된 노드의 프로세스와 포트 정상", ""],
    ["[ ]", "HAProxy", "80 LISTEN CMS login HTTP 200", ""],
    ["[ ]", "학습 배포", "모든 Engine 통지 성공과 Solr 반영", ""],
    ["[ ]", "시뮬레이터", "sessionKey 생성 후 첫 메시지와 왕복 대화", ""],
], [1.4, 3.0, 8.4, 4.4])

g.h1("11 업무 기능 검증")
g.p("인프라와 포트가 정상이어도 업무 기능은 실패할 수 있다. 완료 판정은 반드시 CMS 화면, API 응답 본문, CMS Gateway Engine 로그를 같은 시간대로 확인한다. HTTP 200만으로 성공 처리하지 않는다.")
g.table(["확인", "기능", "검증 방법", "정상 기준", "증적"], [
    ["[ ]", "CMS 로그인", "브라우저에서 /login 접속과 로그인", "화면 표시와 로그인 성공", ""],
    ["[ ]", "봇 조회", "운영 tenant의 봇 목록 조회", "기존 데이터 표시", ""],
    ["[ ]", "대화 시뮬레이션", "대표 시나리오 질의", "정상 응답과 오류 없음", ""],
    ["[ ]", "Solr 검색", "alias를 통한 대표 검색", "기대 문서 반환", ""],
    ["[ ]", "Redis cache", "기능 수행 후 cache key와 TTL 확인", "생성 조회 정상", ""],
    ["[ ]", "학습 배포", "테스트 봇으로 학습과 배포", "상태 완료와 alias 전환", ""],
    ["[ ]", "Scheduler", "예약 배포 또는 허용된 테스트", "중복 실행 없이 성공", ""],
], [1.2, 3.0, 5.1, 4.1, 3.0], font=7.2)
g.h2("학습 배포 필수 확인")
g.p("승인된 정상 테스트 봇으로 CMS에서 학습 배포를 실행한다. CMS가 engine1, engine2, engine3 모두에 업데이트를 통지하고 각 Engine이 학습 완료 로그를 남기는지 확인한다. 1차 실습에서는 8180 방화벽 누락으로 engine2와 engine3 통지가 timeout 되었고, 포트 허용 후 세 노드 모두 성공했다.")
g.code("for h in engine1 engine2 engine3; do\n  echo \"===== $h:8180 =====\"\n  nc -zvw3 \"$h\" 8180\ndone\n\n# 학습 실행 직후 같은 시간대 확인\ngrep -RniE 'Succeed notify Engine Update Dialog|Completed Learning Engine|ERROR|Exception|timeout|No route' \\\n/logs/cms /logs/engine 2>/dev/null | tail -100\n\ncurl -sS 'http://solr1:8983/solr/admin/collections?action=LISTALIASES&wt=json' | python3 -m json.tool")
g.p("정상 기준: CMS에서 배포 완료, 대상 Engine 전부 통지 성공, Engine에서 STANDBY/TEST 등 승인된 profile 학습 완료, Solr collection 또는 alias가 기대 상태다. B2B_BASIC_CODE처럼 source에서도 channel/intent/alias가 없었던 봇은 이관 판정용으로 사용하지 않는다.")

g.h2("시뮬레이터 필수 확인")
g.p("브라우저에서 chat-ui가 열리는 것만으로는 부족하다. Chat UI가 gateway를 통해 세션을 만들고, 응답 JSON에 null이 아닌 sessionKey가 있어야 한다. 200 OK이면서 body가 비어 있으면 애플리케이션 관점에서는 실패다.")
g.code("# 운영에서 승인된 정상 botCode와 channelId로 교체\ncurl -sS --max-time 15 \\\n  -D /tmp/simulator.headers -o /tmp/simulator.body \\\n  -w 'HTTP=%{http_code} SIZE=%{size_download}\\n' \\\n  -X POST -H 'Content-Type: application/json' \\\n  http://gateway/gateway/web/v1/simulator/STANDBY/start \\\n  -d '{\"userKey\":\"<TEST_USER>\",\"callForwardingNumber\":\"<TEST_NUMBER>\",\"channelId\":\"<APPROVED_CHANNEL_ID>\",\"reEntry\":false}'\n\ngrep -oE '\"sessionKey\"[[:space:]]*:[[:space:]]*\"[^\"]+\"' /tmp/simulator.body")
g.p("1차 실습의 TC_CALLBOT 정상 요청은 code 0000, available true, sessionKey gw1_...를 반환했고 브라우저 대화도 동작했다. 반면 B2B_BASIC_CODE는 source와 target 모두 200 + 빈 body와 sessionKey null이어서 기존 봇 데이터 문제로 분류했다. 외부 폰트 CORS 오류는 화면 글꼴 문제이며 세션 생성 실패의 원인이 아니었다.")
g.code("for d in master cms engine gateway chat-ui scheduler; do\n  echo \"===== $d recent errors =====\"\n  grep -Ei 'ERROR|Exception|failed|timeout|refused' \"/logs/$d\"/*.log 2>/dev/null | tail -20\ndone")

g.h1("12 최종 일괄 확인 명령")
g.p("다음 명령은 빠른 상태 요약용이다. 운영 IP와 실제 경로를 먼저 확인한 뒤 실행한다.")
g.code("echo '===== Names ====='\ngetent hosts engine gateway cms master chat-ui engine1 engine2 engine3\n\necho '===== MariaDB MaxScale ====='\nsudo maxctrl list servers\n\necho '===== Redis ====='\n/application/redis/bin/redis-cli -h redis1 -p 7000 cluster info | \\\n  grep -E 'cluster_state|cluster_slots_assigned|cluster_known_nodes'\n\necho '===== ZooKeeper ====='\nfor h in zk1 zk2 zk3; do printf '%-4s ' \"$h\"; printf 'ruok' | nc -w 3 \"$h\" 2181; echo; done\n\necho '===== Solr ====='\ncurl -sS 'http://solr1:8983/solr/admin/collections?action=LIST&wt=json'\ncurl -sS 'http://solr1:8983/solr/admin/collections?action=CLUSTERSTATUS&wt=json' | \\\n  grep -oE '\"state\":\"[^\"]+\"' | sort | uniq -c\n\necho '===== Elasticsearch ====='\ncurl -sS 'http://es1:<ES_HTTP_PORT>/_cat/nodes?v&h=name,ip,version,master'\ncurl -sS 'http://es1:<ES_HTTP_PORT>/_cluster/health?filter_path=status,number_of_nodes,unassigned_shards'\n\necho '===== Applications ====='\nfor d in master cms engine gateway chat-ui scheduler; do /application/$d/status.sh 2>&1; done\nfor h in engine1 engine2 engine3; do nc -zvw3 \"$h\" 8180; done\n\necho '===== Web ====='\ncurl -sS -o /dev/null -w 'CMS_LOGIN_HTTP=%{http_code}\\n' -H 'Host: cms' http://127.0.0.1/login")

g.doc.add_page_break()
g.h1("13 이상 발견 시 판단 순서")
g.table(["증상", "먼저 확인", "다음 조치"], [
    ["No route to host", "방화벽 route 대상 IP", "포트 개방 후 재검증"],
    ["Connection refused", "서비스 프로세스와 LISTEN", "로그 확인 후 서비스 기동"],
    ["STARTED 후 stopped", "로그 디렉터리 data 권한", "소유권과 실제 오류 수정"],
    ["HTTP 500", "애플리케이션 초기화 로그와 의존 서비스", "원인 서비스부터 복구"],
    ["HAProxy 503", "backend 이름 IP 포트 health", "backend 기동 또는 HAProxy 설정 수정"],
    ["브라우저 이름 해석 실패", "PC의 cms chat-ui gateway DNS 또는 hosts", "VIP/LB로 등록 후 DNS cache 제거"],
    ["Master API 403", "인증 없는 직접 호출인지", "서비스 장애와 인증 거부를 구분"],
    ["Solr replica down", "노드 간 8983 통신과 core 로그", "방화벽 수정 후 recovery 확인"],
    ["ES path repo empty", "세 노드 active settings", "모든 노드 설정과 재기동 확인"],
    ["학습 배포 timeout", "engine1~3 이름 8180 방화벽", "모든 Engine 연결 후 재배포"],
    ["시뮬레이터 session null", "Gateway start 응답 body와 channelId", "정상 봇 채널로 재검증"],
    ["HTTP 200 + 빈 body", "응답 크기 Gateway Engine 로그", "성공으로 판정하지 않고 데이터 확인"],
], [4.0, 6.0, 6.3], font=7.9)
g.p("Go 조건: 모든 필수 서비스가 정상이고 source와 target의 승인된 데이터 기준이 일치하며, 정상 테스트 봇의 학습 배포가 모든 Engine에서 완료되고 시뮬레이터가 sessionKey 생성 후 실제 대화를 수행한다.")
g.p("No Go 조건: DB 복제 오류, Redis slot 불완전, ZooKeeper quorum 미구성, Solr down replica, Elasticsearch red, 주요 애플리케이션 반복 종료, CMS login 실패, 학습 배포 실패, 시뮬레이터 session null 중 하나라도 해결되지 않았다.")

g.p("작업 일시, 운영 IP와 VIP, 변경한 설정 파일, 검증 증적, 미해결 예외, Go No Go 결정은 별도 핵심 체크리스트에 기록한다.")

g.save()
print(FINAL)
