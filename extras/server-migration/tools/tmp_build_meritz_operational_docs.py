from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "산출물"
OUT.mkdir(parents=True, exist_ok=True)


def set_font(run, name="맑은 고딕", size=None, bold=None, color=(0, 0, 0)):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rpr.rFonts.set(qn(key), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(*color)


def style_font(style, size, bold=False):
    style.font.name = "맑은 고딕"
    rpr = style._element.get_or_add_rPr()
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rpr.rFonts.set(qn(key), "맑은 고딕")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)


def cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_border(cell, color="D9D9D9"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), color)


def cell_margins(cell, value=100):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for side in ("top", "start", "bottom", "end"):
        node = mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


class Book:
    def __init__(self, title, subtitle, filename, compact=False):
        self.doc = Document()
        self.path = OUT / filename
        sec = self.doc.sections[0]
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11)
        sec.top_margin = Cm(1.7)
        sec.bottom_margin = Cm(1.6)
        sec.left_margin = Cm(1.8)
        sec.right_margin = Cm(1.8)
        sec.footer_distance = Cm(0.7)
        styles = self.doc.styles
        style_font(styles["Normal"], 9.6 if compact else 10.2)
        styles["Normal"].paragraph_format.line_spacing = 1.12 if compact else 1.18
        styles["Normal"].paragraph_format.space_after = Pt(4)
        style_font(styles["Title"], 24, True)
        styles["Title"].paragraph_format.space_after = Pt(11)
        title_ppr = styles["Title"]._element.get_or_add_pPr()
        title_border = title_ppr.find(qn("w:pBdr"))
        if title_border is not None:
            title_ppr.remove(title_border)
        for name, size, before, after in [
            ("Heading 1", 16, 16, 7), ("Heading 2", 13, 11, 5), ("Heading 3", 11, 8, 3)
        ]:
            style_font(styles[name], size, True)
            styles[name].paragraph_format.space_before = Pt(before)
            styles[name].paragraph_format.space_after = Pt(after)
            styles[name].paragraph_format.keep_with_next = True
        if "Code Block" not in styles:
            code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        else:
            code = styles["Code Block"]
        code.font.name = "Consolas"
        code._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
        code.font.size = Pt(7.6 if compact else 8.1)
        code.paragraph_format.left_indent = Cm(0.25)
        code.paragraph_format.right_indent = Cm(0.2)
        code.paragraph_format.space_before = Pt(2)
        code.paragraph_format.space_after = Pt(5)
        code.paragraph_format.line_spacing = 1.0
        p = self.doc.add_paragraph(style="Title")
        p.add_run(title)
        ppr = p._p.get_or_add_pPr()
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)
        p = self.doc.add_paragraph()
        r = p.add_run(subtitle)
        set_font(r, size=11, bold=True, color=(0, 0, 0))
        self.p("작성 기준 2026년 9월 10일 1차 이관 실습 결과")
        self.p("보안 원칙 실제 비밀번호 인증 키 토큰은 문서에 기록하지 않고 자리표시자로 입력한다")
        self.footer()

    def footer(self):
        sec = self.doc.sections[0]
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("메리츠 xen02 to xen01 이관 문서")
        set_font(r, size=8, color=(110, 110, 110))

    def h1(self, text, page=False):
        if page:
            self.doc.add_page_break()
        return self.doc.add_heading(text, 1)

    def h2(self, text):
        return self.doc.add_heading(text, 2)

    def h3(self, text):
        return self.doc.add_heading(text, 3)

    def p(self, text="", lead=None):
        p = self.doc.add_paragraph()
        if lead and text.startswith(lead):
            a = p.add_run(lead)
            a.bold = True
            p.add_run(text[len(lead):])
        else:
            p.add_run(text)
        return p

    def bullets(self, items):
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.add_run(item)

    def numbers(self, items):
        for i, item in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.add_run(f"{i}. {item}")

    def code(self, text):
        p = self.doc.add_paragraph(style="Code Block")
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F3F5F7")
        ppr.append(shd)
        lines = text.strip("\n").splitlines()
        for i, line in enumerate(lines):
            r = p.add_run(line)
            set_font(r, "Consolas", 7.6 if len(line) > 105 else 8.1)
            if i < len(lines) - 1:
                r.add_break()
        return p

    def table(self, headers, rows, widths=None, font=8.5, first_center=True):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Table Grid"
        for i, text in enumerate(headers):
            c = table.rows[0].cells[i]
            cell_shading(c, "1F4E78")
            cell_border(c)
            cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(text))
            set_font(r, size=8.6, bold=True, color=(255, 255, 255))
            if widths:
                c.width = Cm(widths[i])
        trpr = table.rows[0]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:tblHeader"))
        for ridx, row in enumerate(rows):
            cells = table.add_row().cells
            for i, text in enumerate(row):
                c = cells[i]
                if ridx % 2:
                    cell_shading(c, "F6F9FC")
                cell_border(c)
                cell_margins(c)
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if first_center and i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(str(text))
                set_font(r, size=font)
                if widths:
                    c.width = Cm(widths[i])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)
        return table

    def save(self):
        self.doc.save(self.path)
        return self.path


HOSTS = """192.168.1.81 redis1 zk1 solr1 es1 mariadb1 engine1 gateway1 cms1 master1 chat-ui1 scheduler1 mariadb engine gateway cms master chat-ui
192.168.1.82 redis2 zk2 solr2 es2 mariadb2 engine2 gateway2 cms2 master2 chat-ui2
192.168.1.83 redis3 zk3 solr3 es3 mariadb3 engine3 gateway3 cms3 master3 chat-ui3"""


def build_runbook():
    b = Book(
        "메리츠 xen02 to xen01 운영 이관 런북",
        "승인된 작업자가 순서대로 실행하고 각 단계의 정상 기준을 확인하는 절차",
        "메리츠_xen02_xen01_운영용_런북.docx",
    )
    b.h1("문서 목적과 운영 원칙")
    b.p("이 런북은 xen02 3대의 메리츠 환경을 xen01 3대로 이관할 때 사용하는 실행 절차다. 1차 실습에서 확인한 구성과 오류를 반영했으며, 각 단계는 명령 실행 후 정상 기준을 확인한 다음 진행한다.")
    b.bullets([
        "source는 xen02.01부터 xen02.03까지이며 IP는 192.168.1.101부터 192.168.1.103까지다.",
        "target은 xen01.01부터 xen01.03까지이며 IP는 192.168.1.81부터 192.168.1.83까지다.",
        "기존 xen01.02의 Elasticsearch 9.1.1은 9200과 9300에서 유지한다.",
        "메리츠 Elasticsearch 7.8.0은 9210과 9310에서 별도 클러스터로 운영한다.",
        "Redis cache는 원칙적으로 재생성한다. 기존 key 이관은 업무 승인 후에만 수행한다.",
        "Solr 데이터는 Collections API BACKUP과 RESTORE로 이관한다.",
        "운영 전환 시 source writer 중지 시점과 마지막 동기화 시점을 기록한다.",
    ])
    b.h2("최종 구성")
    b.table(["영역", "구성", "포트", "정상 기준"], [
        ["MariaDB", "1 master 2 slave MaxScale", "13306 3306", "복제 정상 MaxScale Running"],
        ["Redis", "3 master 3 replica", "7000 7001 17000 17001", "16384 slots covered"],
        ["ZooKeeper", "1 leader 2 follower", "2181 2888 3888", "3대 imok"],
        ["SolrCloud", "3 nodes", "8983", "6 collections all active"],
        ["Elasticsearch", "메리츠 7.8.0 3 nodes", "9210 9310", "green"],
        ["기존 Elasticsearch", "xen01.02 9.1.1", "9200 9300", "기존 cluster 유지"],
        ["Applications", "공통 5개 scheduler 1개", "8081 8180 8280 8380 8480 8580", "process와 HTTP 정상"],
        ["HAProxy", "xen01.01", "80", "cms login HTTP 200"],
    ], widths=[3.0, 5.2, 3.4, 5.0])

    b.h1("사전 준비")
    b.h2("작업 전 승인과 기록")
    b.bullets([
        "변경 시간과 담당자, source와 target IP, 승인 번호를 기록한다.",
        "source 쓰기 중지 여부와 허용 중단 시간을 확인한다.",
        "복구 기준은 source 서비스 재개와 target 서비스 중지로 정의한다.",
        "비밀번호와 인증 키는 쉘 기록이나 문서에 직접 남기지 않는다.",
    ])
    b.h2("서버 식별")
    b.code("hostname\nhostname -i\ncat /etc/os-release\ndate '+%F %T %Z'\ndf -hT / /data /application /tmp 2>/dev/null")
    b.p("정상 기준: 현재 접속한 서버의 hostname과 IP가 작업표의 source 또는 target과 일치한다.")
    b.h2("호스트 이름 등록")
    b.code(HOSTS)
    b.code("getent hosts mariadb redis1 redis2 redis3 zk1 zk2 zk3 solr1 solr2 solr3 es1 es2 es3 master engine3 scheduler1")
    b.p("정상 기준: 모든 이름이 target IP를 반환한다. source와 target이 동시에 켜져 있을 때 같은 이름을 양쪽 IP에 중복 등록하지 않는다.")
    b.h2("필수 포트와 방화벽")
    b.table(["서비스", "포트"], [
        ["HAProxy", "80"], ["MaxScale", "3306"], ["MariaDB", "13306"],
        ["Redis", "7000 7001 17000 17001"], ["ZooKeeper", "2181 2888 3888"],
        ["Solr", "8983"], ["Elasticsearch 메리츠", "9210 9310"],
        ["Applications", "8081 8180 8280 8380 8480 8580"],
    ], widths=[6.0, 9.0])
    b.code("sudo firewall-cmd --zone=public --list-all\nsudo ss -lntp")

    b.h1("MariaDB와 MaxScale")
    b.h2("복제 상태 확인")
    b.code("/mariadb/bin/mariadb -S /mariadb/run/mariadb.sock -e \"SHOW SLAVE STATUS\\G\" | \\\ngrep -E 'Slave_IO_Running|Slave_SQL_Running|Seconds_Behind_Master|Last_SQL_Error|Gtid_IO_Pos'")
    b.p("정상 기준: replica 두 대의 IO와 SQL thread가 Yes이고 Seconds Behind Master가 0이며 Last SQL Error가 비어 있다.")
    b.h2("MaxScale 서비스 계정")
    b.code("/mariadb/bin/mariadb -S /mariadb/run/mariadb.sock -e \\\n+\"SELECT User,Host,plugin FROM mysql.user WHERE User='mxs_service';\"\n\nsudo maxctrl alter service MariaDB-Service \\\n+user=mxs_service password='<MXS_SERVICE_PASSWORD>'")
    b.p("계정은 사용자 이름과 접속 Host의 조합이다. mxs_service 계정의 Host는 MaxScale 접속 IP와 정확히 일치해야 한다.")
    b.h2("최종 확인")
    b.code("sudo maxctrl list servers\nnc -zvw3 mariadb 3306")
    b.p("정상 기준: mariadb1은 Master Running, mariadb2와 mariadb3은 Slave Running이다. GTID 차이가 있으면 전환 전에 다시 복제 상태를 확인한다.")

    b.h1("Redis Cluster")
    b.h2("runtime 배포")
    b.code("# source에서 데이터 디렉터리를 제외하고 생성\nsudo tar -czpf /tmp/redis_runtime.tar.gz -C /application --exclude=redis/db redis\n\n# target에서 압축 해제\nsudo mkdir -p /application\nsudo tar -xzpf /tmp/redis_runtime.tar.gz -C /application\nsudo install -d -o admin -g admin -m 750 /logs/redis")
    b.h2("노드별 설정")
    b.code("# 각 서버의 7000.conf와 7001.conf\nport 7000 또는 7001\ndir /application/redis/db/7000 또는 7001\ncluster-enabled yes\ncluster-announce-ip redis1 또는 redis2 또는 redis3\ncluster-announce-port 7000 또는 7001\ncluster-announce-bus-port 17000 또는 17001")
    b.p("각 target 서버에서 announce IP 별칭이 해당 서버 자신을 가리키는지 확인한다.")
    b.h2("기동과 통신 확인")
    b.code("/application/redis/start.sh\n/application/redis/status.sh\n\nfor h in redis1 redis2 redis3; do\n  for p in 7000 7001 17000 17001; do nc -zvw2 \"$h\" \"$p\"; done\ndone")
    b.h2("새 클러스터 생성")
    b.code("/application/redis/bin/redis-cli --cluster create \\\n+192.168.1.81:7000 192.168.1.82:7000 192.168.1.83:7000 \\\n+192.168.1.81:7001 192.168.1.82:7001 192.168.1.83:7001 \\\n+--cluster-replicas 1")
    b.p("대상 노드가 비어 있고 기존 cluster state가 없을 때만 실행한다. 확인 질문에는 슬롯과 replica 배치를 검토한 뒤 yes를 입력한다.")
    b.h2("최종 확인")
    b.code("/application/redis/bin/redis-cli --cluster check redis1:7000\n/application/redis/bin/redis-cli -h redis1 -p 7000 cluster info")
    b.p("정상 기준: 6 nodes, cluster size 3, cluster state ok, 16384 slots covered, open slot 없음.")

    b.h1("ZooKeeper")
    b.h2("runtime 배포와 myid")
    b.code("# source\nsudo tar -czpf /tmp/zk_runtime.tar.gz -C /application/solr --exclude=zk/data --exclude=zk/datalog zk\n\n# target\nsudo mkdir -p /application/solr\nsudo tar -xzpf /tmp/zk_runtime.tar.gz -C /application/solr\nsudo install -d -o admin -g admin -m 750 /logs/solr/zk\nsudo install -d -o admin -g admin -m 750 /application/solr/zk/data /application/solr/zk/datalog\nprintf '1\\n' | sudo tee /application/solr/zk/data/myid  # zk1\nprintf '2\\n' | sudo tee /application/solr/zk/data/myid  # zk2\nprintf '3\\n' | sudo tee /application/solr/zk/data/myid  # zk3")
    b.h2("설정")
    b.code("dataDir=/application/solr/zk/data\ndataLogDir=/application/solr/zk/datalog\nclientPort=2181\nserver.1=zk1:2888:3888\nserver.2=zk2:2888:3888\nserver.3=zk3:2888:3888")
    b.h2("기동과 확인")
    b.code("/application/solr/zk/bin/zkServer.sh start\n\nfor h in zk1 zk2 zk3; do\n  printf 'ruok' | nc -w 3 \"$h\" 2181; echo\n  printf 'srvr' | nc -w 3 \"$h\" 2181 | grep -E 'Mode|Node count'\ndone")
    b.p("정상 기준: 세 노드 모두 imok, leader 1대, follower 2대, Node count 동일.")

    b.h1("SolrCloud")
    b.h2("전용 ZooKeeper 경로")
    b.p("source와 target Solr가 같은 ZooKeeper ensemble을 동시에 사용할 수 있으므로 target은 전용 chroot인 solr xen01을 사용한다.")
    b.code("# 노드별 solr.in.sh\nSOLR_HOST=\"solr1\"  # 두 번째와 세 번째 노드는 solr2 solr3\nSOLR_PORT=\"8983\"\nZK_HOST=\"zk1:2181,zk2:2181,zk3:2181/solr-xen01\"")
    b.h2("configset 이관")
    b.code("# source\n/application/solr/solr/bin/solr zk downconfig \\\n+-n chat-base-config -d /tmp/solr_config_export/chat-base-config \\\n+-z zk1:2181,zk2:2181,zk3:2181/solr\n\n# target\n/application/solr/solr/bin/solr zk upconfig \\\n+-n chat-base-config -d /tmp/solr_config_import/chat-base-config \\\n+-z zk1:2181,zk2:2181,zk3:2181/solr-xen01")
    b.h2("Solr 기동")
    b.code("/application/solr/solr/bin/solr start -cloud -p 8983 \\\n+-z zk1:2181,zk2:2181,zk3:2181/solr-xen01")
    b.h2("collection 백업과 복구")
    b.p("공유 NFS 경로가 source와 target의 모든 Solr 노드에 동일하게 mount되어 있어야 한다.")
    b.code("# source 백업 예시\ncurl -sS --max-time 300 \\\n+'http://127.0.0.1:8983/solr/admin/collections?action=BACKUP&collection=TC_CALLBOT_B&name=TC_CALLBOT_B_20260910&location=/data/solr_backup/TC_CALLBOT_B&wt=json'\n\n# target 복구 예시\ncurl -sS --max-time 300 \\\n+'http://127.0.0.1:8983/solr/admin/collections?action=RESTORE&collection=TC_CALLBOT_B&name=TC_CALLBOT_B_20260910&location=/data/solr_backup/TC_CALLBOT_B&wt=json'")
    b.p("같은 방식으로 여섯 collection을 모두 처리한다. TC CALLBOT LOG, TC CALLBOT B, tc 01 A, TC CALLBOT TEST, tc 01 B, TC CALLBOT A가 대상이다.")
    b.h2("Solr 검증")
    b.code("curl -sS 'http://solr1:8983/solr/admin/collections?action=LIST&wt=json'\ncurl -sS 'http://solr1:8983/solr/admin/collections?action=LISTALIASES&wt=json'\ncurl -sS 'http://solr1:8983/solr/admin/collections?action=CLUSTERSTATUS&wt=json' | grep -oE '\"state\":\"[^\"]+\"' | sort | uniq -c")
    b.p("정상 기준: 6 collections, alias TC CALLBOT과 tc 01 존재, 모든 replica active, source와 target의 문서 수 일치.")

    b.h1("공유 NFS 백업 경로")
    b.h2("NFS server")
    b.code("sudo install -d -o admin -g admin -m 770 /data/solr_backup\nsudo vi /etc/exports.d/solr-migration.exports\nsudo systemctl enable --now nfs-server\nsudo exportfs -rav\nsudo firewall-cmd --permanent --add-service=nfs\nsudo firewall-cmd --permanent --add-service=mountd\nsudo firewall-cmd --permanent --add-service=rpc-bind\nsudo firewall-cmd --reload")
    b.h2("NFS client")
    b.code("sudo install -d -o admin -g admin -m 770 /data/solr_backup\nsudo mount -t nfs 192.168.1.101:/data/solr_backup /data/solr_backup\nfindmnt /data/solr_backup\ntouch /data/solr_backup/write-test-$(hostname -i | tr . -)")
    b.p("root squash가 설정되어 있으므로 NFS 파일은 admin 권한으로 작성하고, sudo로 읽을 때 Permission denied가 나면 먼저 /tmp로 복사한 뒤 처리한다.")

    b.h1("Elasticsearch 7 8 0")
    b.h2("기존 9 1 1 보호")
    b.p("xen01.02의 기존 9.1.1은 9200과 9300에서 유지한다. 메리츠 7.8.0은 9210과 9310을 사용하므로 두 데이터 경로와 로그 경로가 절대 겹치면 안 된다.")
    b.h2("runtime과 노드별 설정")
    b.code("# source runtime archive\nsudo tar -czpf /tmp/elk78_runtime.tar.gz -C /application --exclude=elk/data --exclude=elk/tmp elk\n\n# target elasticsearch.yml 핵심\ncluster.name: \"trusted-context-es\"\nnode.name: \"es1\"\npath.data: \"/application/elk-meritz/data\"\npath.logs: \"/logs/elk-meritz\"\nnetwork.host: 192.168.1.81\nhttp.port: 9210\ntransport.port: 9310\ndiscovery.seed_hosts: [\"es1:9310\", \"es2:9310\", \"es3:9310\"]\npath.repo: [\"/data/solr_backup\"]")
    b.p("es2와 es3은 node name과 network host만 해당 노드 값으로 변경한다. jvm.options의 HeapDumpPath, ErrorFile, GC log도 /logs/elk-meritz로 변경한다.")
    b.h2("OS 설정과 기동")
    b.code("sudo sysctl -w vm.max_map_count=262144\nsudo install -d -o admin -g admin -m 750 /logs/elk-meritz /application/elk-meritz/data /application/elk-meritz/run\nES_PATH_CONF=/application/elk/config /application/elk/bin/elasticsearch -d -p /application/elk-meritz/run/elasticsearch.pid")
    b.h2("source snapshot")
    b.code("curl -sS -X PUT -H 'Content-Type: application/json' \\\n+http://127.0.0.1:9200/_snapshot/meritz_migration_repo \\\n+-d '{\"type\":\"fs\",\"settings\":{\"location\":\"/data/solr_backup/elasticsearch\"}}'\n\ncurl -sS -X PUT -H 'Content-Type: application/json' \\\n+'http://127.0.0.1:9200/_snapshot/meritz_migration_repo/meritz_20260910?wait_for_completion=true' \\\n+-d '{\"indices\":\"*\",\"include_global_state\":true}'")
    b.h2("target repository와 restore")
    b.code("curl -sS -X PUT -H 'Content-Type: application/json' \\\n+http://es1:9210/_snapshot/meritz_migration_repo \\\n+-d '{\"type\":\"fs\",\"settings\":{\"location\":\"/data/solr_backup/elasticsearch\",\"readonly\":true}}'\n\ncurl -sS -X POST http://es1:9210/_snapshot/meritz_migration_repo/_verify?pretty\n\ncurl -sS -X POST -H 'Content-Type: application/json' \\\n+'http://es1:9210/_snapshot/meritz_migration_repo/meritz_20260910/_restore?wait_for_completion=true' \\\n+-d '{\"indices\":\"*\",\"include_global_state\":true}'")
    b.h2("Elasticsearch 검증")
    b.code("curl -sS 'http://es1:9210/_cat/nodes?v&h=name,ip,role,master,version'\ncurl -sS 'http://es1:9210/_cluster/health?pretty'\ncurl -sS 'http://es1:9210/_cat/indices?expand_wildcards=all&v&h=health,status,index,docs.count,store.size&s=index'")
    b.p("정상 기준: 3 nodes, version 7.8.0, cluster green, 6 indices와 문서 수가 source와 일치한다.")

    b.h1("애플리케이션과 HAProxy")
    b.h2("배치 구조")
    b.p("gateway, engine, master, cms, chat ui는 세 target 서버에 공통 배치한다. scheduler는 중복 실행을 막기 위해 승인된 한 노드에서만 기동한다.")
    b.h2("설정 변경")
    b.code("# gateway\nelasticsearch.url: http://es1:9210\nelasticsearch.port: 9210\n\n# master cms scheduler\nelasticsearch.host1: es1\nelasticsearch.host2: es2\nelasticsearch.host3: es3\nelasticsearch.port: 9210\n\n# 기존 의존성\nMariaDB: mariadb:3306\nRedis: redis1:7000,redis2:7000,redis3:7000\nSolr: solr1:8983,solr2:8983,solr3:8983\nMaster: http://master\nScheduler: http://scheduler1:8580/scheduler")
    b.h2("로그 디렉터리")
    b.code("sudo install -d -o admin -g admin -m 750 \\\n+/logs/master /logs/cms /logs/engine /logs/gateway /logs/chat-ui /logs/scheduler")
    b.h2("기동 순서")
    b.numbers([
        "MariaDB와 MaxScale 상태를 확인한다.",
        "Redis, ZooKeeper, SolrCloud, Elasticsearch를 확인한다.",
        "Master를 시작하고 8380 LISTEN을 확인한다.",
        "CMS를 시작하고 8280 LISTEN과 로그를 확인한다.",
        "Engine, Gateway, Chat UI를 시작한다.",
        "Scheduler는 승인된 한 노드에서 마지막에 시작한다.",
    ])
    b.code("/application/master/start.sh\n/application/cms/start.sh\n/application/engine/start.sh\n/application/gateway/start.sh\n/application/chat-ui/start.sh\n/application/scheduler/start.sh")
    b.h2("HAProxy와 웹")
    b.code("sudo haproxy -c -f /etc/haproxy/haproxy.cfg\nsudo systemctl enable --now haproxy\nsudo systemctl restart haproxy\nsudo ss -lntp | grep ':80\\b'\n\ncurl -sS -o /dev/null -w 'HTTP_CODE=%{http_code}\\n' -H 'Host: cms' http://127.0.0.1/login")
    b.p("정상 기준: HAProxy 80 LISTEN, cms login HTTP 200. 사용자 PC hosts에는 cms와 meritz.easycms를 192.168.1.81로 한 번만 등록한다.")

    b.h1("최종 검증과 전환")
    b.h2("의존 서비스")
    b.code("nc -zvw2 mariadb 3306\n/application/redis/bin/redis-cli -h redis1 -p 7000 cluster info | grep -E 'cluster_state|cluster_slots_assigned|cluster_known_nodes'\nfor h in zk1 zk2 zk3; do printf 'ruok' | nc -w 3 \"$h\" 2181; echo; done\ncurl -sS 'http://solr1:8983/solr/admin/info/system?wt=json' | grep -E '\"mode\"|\"zkHost\"|\"solr-spec-version\"'\ncurl -sS http://es1:9210/ | grep -E '\"name\"|\"cluster_name\"|\"number\"'")
    b.h2("애플리케이션")
    b.code("for d in master cms engine gateway chat-ui scheduler; do\n  echo \"===== $d =====\"\n  /application/$d/status.sh 2>&1\ndone\nsudo ss -lntp | grep -E ':(8081|8180|8280|8380|8480|8580)\\b'")
    b.h2("업무 기능")
    b.bullets([
        "CMS 로그인과 사용자 인증",
        "bot 목록과 상세 조회",
        "대화 시뮬레이터",
        "Solr 검색과 alias 조회",
        "Redis cache 생성과 조회",
        "학습과 배포 및 예약 배치",
        "애플리케이션 오류 로그와 HTTP 5xx 부재",
    ])
    b.p("B2B BASIC CODE 학습 실패는 source에서도 동일하게 발생했고 해당 Solr alias가 없었다. 이 항목은 이관 실패가 아니라 기존 업무 데이터 또는 설정 문제로 분리한다.")
    b.h2("복구 판단")
    b.p("핵심 기능이 승인 시간 내 복구되지 않거나 데이터 건수가 source와 다르면 target writer와 Scheduler를 중지하고 사용자 PC와 DNS 또는 VIP를 source로 되돌린다. source 서비스 재개 후 원인을 분석한다.")
    return b.save()


def build_checklist():
    b = Book(
        "메리츠 xen02 to xen01 핵심 체크리스트",
        "작업자가 단계별 완료 여부와 결과를 기록하는 현장 점검표",
        "메리츠_xen02_xen01_핵심_체크리스트.docx",
        compact=True,
    )
    b.h1("사용 방법")
    b.p("각 항목은 확인 명령의 실제 결과를 본 뒤 완료 표시한다. 포트가 열렸다는 사실만으로 정상 판정하지 않으며 cluster state, 데이터 건수, HTTP 응답, 업무 기능을 함께 확인한다.")
    b.table(["기록 항목", "값"], [["작업 일시", ""], ["담당자", ""], ["변경 승인", ""], ["source 쓰기 중지", "예 아니오"], ["전환 시작", ""], ["전환 완료", ""]], widths=[5.0, 10.0])
    sections = [
        ("사전 점검", [
            ("현재 서버 hostname과 IP 확인", "hostname; hostname -i", "작업표와 일치"),
            ("디스크 여유 확인", "df -hT / /data /application /tmp", "백업 크기 이상 여유"),
            ("target hosts 확인", "getent hosts mariadb redis1 zk1 solr1 es1 master", "모두 target IP"),
            ("방화벽 정책 확인", "firewall-cmd --zone=public --list-all", "필수 포트 허용"),
            ("source와 target 구분 기록", "date '+%F %T %Z'", "시간과 대상 기록"),
        ]),
        ("MariaDB와 MaxScale", [
            ("replica IO thread", "SHOW SLAVE STATUS", "Yes"),
            ("replica SQL thread", "SHOW SLAVE STATUS", "Yes"),
            ("복제 지연", "Seconds_Behind_Master", "0"),
            ("복제 오류", "Last_SQL_Error", "빈 값"),
            ("MaxScale backend", "maxctrl list servers", "1 Master 2 Slave Running"),
            ("MaxScale listener", "nc -zvw2 mariadb 3306", "Connected"),
        ]),
        ("Redis", [
            ("6개 프로세스", "status.sh", "모두 ready yes"),
            ("client와 bus 통신", "7000 7001 17000 17001", "3대 모두 Connected"),
            ("cluster state", "redis-cli cluster info", "ok"),
            ("known nodes", "redis-cli cluster info", "6"),
            ("slot coverage", "redis-cli --cluster check", "16384 covered"),
            ("replica 구조", "cluster nodes", "3 master 3 replica"),
        ]),
        ("ZooKeeper", [
            ("myid", "cat data/myid", "zk1 1 zk2 2 zk3 3"),
            ("client 응답", "ruok", "3대 imok"),
            ("quorum", "srvr", "leader 1 follower 2"),
            ("데이터 일관성", "srvr Node count", "3대 동일"),
        ]),
        ("SolrCloud", [
            ("Solr 프로세스", "curl info system", "3대 HTTP 정상"),
            ("ZooKeeper chroot", "zkHost", "solr xen01"),
            ("live nodes", "CLUSTERSTATUS", "solr1 solr2 solr3"),
            ("collections", "LIST", "6개"),
            ("replicas", "CLUSTERSTATUS", "22 active 0 down"),
            ("aliases", "LISTALIASES", "TC CALLBOT과 tc 01"),
            ("문서 수", "select rows 0", "source와 일치"),
        ]),
        ("Elasticsearch", [
            ("기존 9 1 1", "xen01.02 9200", "기존 cluster 응답"),
            ("메리츠 노드", "cat nodes 9210", "es1 es2 es3 7.8.0"),
            ("cluster health", "cluster health", "green"),
            ("snapshot repository", "snapshot verify POST", "3 nodes"),
            ("indices", "cat indices", "6개 source와 동일"),
            ("path repo", "nodes settings", "3대 모두 동일"),
        ]),
        ("애플리케이션과 웹", [
            ("의존 서비스 접속", "DB Redis ZK Solr ES", "모두 정상"),
            ("로그 디렉터리", "/logs 각 서비스", "admin 쓰기 가능"),
            ("Master", "status와 8380", "running LISTEN"),
            ("CMS", "status와 8280", "running LISTEN"),
            ("Engine Gateway Chat UI", "status와 ports", "모두 running"),
            ("Scheduler", "전체 서버 확인", "승인 노드 1개만 running"),
            ("HAProxy", "haproxy check와 80", "valid LISTEN"),
            ("CMS 로그인", "Host cms GET login", "HTTP 200"),
            ("Windows hosts", "ping cms", "192.168.1.81"),
        ]),
        ("업무 검증과 종료", [
            ("로그인과 권한", "브라우저", "정상"),
            ("bot 조회", "CMS", "정상"),
            ("대화 시뮬레이터", "Chat UI", "정상"),
            ("검색", "Solr alias query", "정상"),
            ("cache", "Redis", "생성 조회 정상"),
            ("학습 배포", "CMS와 batch log", "업무 기준 충족"),
            ("오류 로그", "최근 ERROR Exception", "신규 치명 오류 없음"),
            ("source target 최종 건수", "DB Solr ES", "일치"),
            ("전환 또는 복구 결정", "승인자", "결정 기록"),
        ]),
    ]
    for idx, (title, rows) in enumerate(sections):
        b.h1(title, page=(idx > 0 and idx % 2 == 0))
        b.table(["확인", "점검 항목", "확인 방법", "정상 기준", "결과와 시간"], [["[ ]", *r, ""] for r in rows], widths=[1.0, 4.2, 4.2, 4.2, 3.4], font=8.0)
    b.h1("최종 승인")
    b.table(["구분", "판정과 서명"], [["전환 승인", ""], ["복구 승인", ""], ["미해결 항목", ""], ["완료 시각", ""]], widths=[4.5, 11.5])
    return b.save()


def build_troubleshooting():
    b = Book(
        "메리츠 xen02 to xen01 트러블슈팅 부록",
        "1차 실습에서 발생한 오류의 의미와 확인 순서와 해결 방향",
        "메리츠_xen02_xen01_트러블슈팅_부록.docx",
    )
    b.h1("사용 원칙")
    b.bullets([
        "오류 문구를 지우거나 같은 명령을 반복하기 전에 process, port, log, config, data state를 확인한다.",
        "source와 target hostname과 IP를 먼저 확인한다.",
        "삭제, FLUSH, RESET, CREATE OR REPLACE, 복제 변경은 대상 서버를 다시 확인한 뒤 수행한다.",
        "정상 메시지 하나만 믿지 않는다. 예를 들어 STARTED 후 JVM이 종료될 수 있으므로 port와 log를 함께 본다.",
    ])
    b.h2("공통 진단 명령")
    b.code("hostname; hostname -i; date '+%F %T %Z'\nps -ef | grep -E '[j]ava|[r]edis|[h]aproxy'\nsudo ss -lntp\nsudo firewall-cmd --zone=public --list-all\ngetent hosts mariadb redis1 zk1 solr1 es1 master")

    issues = [
        ("MariaDB 1045 Access denied", "비밀번호 불일치 또는 User와 Host 조합 불일치", "SELECT User,Host,plugin FROM mysql.user WHERE User='mxs_service'; MaxScale 접속 IP 확인", "정확한 Host 계정을 만들거나 암호를 맞추고 MaxScale runtime 계정도 같은 값으로 변경한다."),
        ("MariaDB 1130 Host not allowed", "client IP에 대응하는 계정 행이 없음", "mysql.user의 Host 목록과 실제 접속 출발지 확인", "mxs_service at MaxScale IP 계정을 생성하고 필요한 최소 권한을 부여한다."),
        ("ALTER USER 1396 또는 SET PASSWORD 1133", "계정 행이 조회 결과와 실제 권한 테이블에서 불일치하거나 복제 DDL 순서가 꼬임", "SHOW CREATE USER, mysql.global_priv, SHOW SLAVE STATUS 확인", "replica에서 임의 반복하지 말고 source 계정 상태와 binary log 순서를 확인한다."),
        ("sql log bin은 SESSION 변수", "SET GLOBAL을 사용함", "SHOW VARIABLES LIKE sql_log_bin", "동일한 MariaDB client 세션에서 SET SESSION sql_log_bin OFF 후 필요한 로컬 조치를 수행하고 다시 ON으로 돌린다."),
        ("GTID out of order", "replica 로컬 DDL의 GTID domain과 source GTID sequence 충돌", "SHOW SLAVE STATUS의 Last SQL Error와 Gtid IO Pos 확인", "임의 skip 전에 데이터 일관성을 평가하고 승인된 GTID 복구 절차를 적용한다."),
        ("repl user ALTER USER 실패", "source와 replica의 계정 상태 차이", "User Host 행과 replication SQL error 확인", "모든 노드 계정 상태를 맞춘 뒤 SQL thread를 재개한다."),
        ("Redis No route to host 또는 timeout", "방화벽 또는 네트워크 경로 차단", "nc로 7000 7001 17000 17001 전체 조합 확인", "각 노드 방화벽에서 client와 cluster bus 포트를 허용하고 다시 검사한다."),
        ("Redis node is not empty", "key 또는 기존 cluster state가 남아 있음", "DBSIZE와 CLUSTER NODES와 nodes config 파일 확인", "target의 보존 필요성을 확인한 후 빈 노드로 정리한다. source에는 RESET이나 FLUSH를 실행하지 않는다."),
        ("Redis Invalid node address", "hostname 형식이나 구버전 redis cli 주소 처리 문제", "getent hosts와 redis cli version 확인", "cluster create는 target IP 주소를 사용하고 announce 주소도 검증한다."),
        ("Redis MOVED 또는 NOKEY", "cluster aware가 아닌 SCAN MIGRATE 반복 또는 TTL 만료", "CLUSTER KEYSLOT과 owner node 확인", "cluster aware 방식으로 이관하거나 cache는 서비스가 재생성하게 한다."),
        ("ZooKeeper log Permission denied", "/logs/solr/zk가 없거나 admin 쓰기 불가", "ls -ld와 sudo -u admin test -w 확인", "install -d로 admin 소유 750 디렉터리를 만든다."),
        ("ZooKeeper Unable to create datalog version 2", "dataDir 또는 dataLogDir 소유권 오류", "설정 경로와 ls -ld와 write test 확인", "data와 datalog를 admin 소유로 바꾸고 최소 권한을 적용한다."),
        ("ZooKeeper STARTED 후 status 실패", "JVM이 시작 직후 종료됨", "QuorumPeerMain process, 2181 LISTEN, service log 확인", "로그의 실제 종료 원인을 수정한 뒤 재기동한다."),
        ("nc command not found", "nmap ncat 미설치", "command -v nc", "nmap-ncat을 설치하거나 Bash의 dev tcp를 사용한다. ZooKeeper 서비스 자체 장애는 아니다."),
        ("Solr runtime 파일 없음", "xen01.02와 xen01.03에 runtime 미배포", "ls solr bin과 solr in sh", "archive를 배포하고 노드별 SOLR HOST를 수정한다."),
        ("Solr port 8983 already used", "기존 Solr process가 실행 중", "ss -lntp와 solr status", "기존 process와 ZK chroot를 확인하고 중복 start하지 않는다."),
        ("Solr HTTP 500 CoreContainer unavailable", "중복 SOLR HOST, 잘못된 ZK chroot, stale core 또는 초기화 실패", "solr log, solr in sh, live nodes, core directories 확인", "solr1 solr2 solr3를 고유하게 설정하고 target 전용 chroot를 사용한다."),
        ("Solr replica down과 query timeout", "노드 간 8983 통신 실패 또는 복사된 stale core", "CLUSTERSTATUS, curl 각 노드, firewall, core properties 확인", "통신을 열고 target RESTORE가 생성한 replica만 남겨 재확인한다."),
        ("NFS Permission denied", "root squash로 client root가 익명 사용자로 매핑됨", "exportfs -v와 파일 owner 확인", "admin으로 공유 경로에 복사한다. 필요하면 /tmp를 경유하고 로컬에서 sudo로 압축 해제한다."),
        ("Elasticsearch GC log path 오류", "jvm options에 source 로그 경로가 남음", "grep logs elk config jvm options", "HeapDumpPath, ErrorFile, GC log와 path logs를 elk meritz 전용 경로로 맞춘다."),
        ("Elasticsearch failed to obtain node lock", "data 경로 권한 오류 또는 같은 data를 쓰는 process 존재", "ps, lsof, data owner와 node lock 확인", "중복 process를 종료하고 노드마다 고유한 writable data path를 사용한다."),
        ("Elasticsearch vm max map count", "OS 값이 262144보다 낮음", "sysctl vm.max_map_count", "sysctl로 262144 이상 적용하고 영구 설정 파일에 기록한다."),
        ("Elasticsearch localhost connection refused", "network host가 node IP에만 bind됨", "ss -lntp에서 실제 bind 주소 확인", "127.0.0.1 대신 해당 node IP 또는 es1 이름으로 호출한다."),
        ("Elasticsearch path repo empty", "일부 노드에 path repo 미설정 또는 변경 후 미재기동", "nodes settings API로 세 노드를 각각 확인", "세 노드 모두 같은 path.repo를 설정하고 rolling restart한 뒤 repository를 등록한다."),
        ("Elasticsearch repository verify 호출 오류", "verify endpoint에 PUT 또는 GET 사용", "HTTP method와 응답 확인", "repository 생성은 PUT, verify는 POST를 사용한다."),
        ("애플리케이션 log Permission denied", "/logs 하위 서비스 디렉터리 없음", "ls -ld와 start log 확인", "서비스별 로그 디렉터리를 admin 소유로 만든 후 재기동한다."),
        ("CMS started 후 stopped", "Master 미기동 또는 master 이름 해석 실패", "CMS log, getent hosts master, nc master 8380", "Master를 먼저 기동하고 master 별칭을 target 주소에 등록한 뒤 CMS를 시작한다."),
        ("Master API 403 Auth key error", "인증 없는 관리 API를 curl로 직접 호출", "실제 CMS 통신과 source 동일 호출 결과 비교", "403 하나로 장애 판정하지 않는다. CMS가 인증 정보를 포함해 정상 통신하는지 확인한다."),
        ("HAProxy 503 Service Unavailable", "backend process down, 이름 해석 실패 또는 방화벽 차단", "HAProxy config check, backend ports, getent hosts, HAProxy log 확인", "backend를 정상화하고 host와 firewall을 수정한 뒤 HAProxy를 reload한다."),
        ("ping 성공 HTTP timeout", "ICMP는 허용됐지만 TCP 80이 차단됨", "curl verbose와 firewall services 확인", "target firewall에 HTTP 서비스를 허용한다."),
        ("학습배포 B2B BASIC CODE 실패", "해당 bot의 Solr alias 또는 업무 데이터가 없음", "LISTALIASES와 source target CMS log 비교", "source에서도 동일하면 이관 장애로 보지 않고 업무 데이터 담당자에게 전달한다."),
    ]
    for title, cause, check, action in issues:
        heading = b.h2(title)
        cause_p = b.p("원인: " + cause, "원인:")
        check_p = b.p("확인: " + check, "확인:")
        action_p = b.p("조치: " + action, "조치:")
        heading.paragraph_format.keep_with_next = True
        cause_p.paragraph_format.keep_with_next = True
        check_p.paragraph_format.keep_with_next = True
        for paragraph in (heading, cause_p, check_p, action_p):
            paragraph.paragraph_format.keep_together = True

    b.h1("명령 입력 실수", page=True)
    b.table(["증상", "의미", "예방과 복구"], [
        ["command not found에 상태 출력 문구 표시", "출력 결과를 명령으로 다시 붙여 넣음", "프롬프트와 출력은 제외하고 명령 줄만 복사"],
        ["bash admin at xen01 command not found", "프롬프트까지 복사", "달러 기호 뒤 명령만 입력"],
        ["bracketed paste 제어문자", "터미널 붙여넣기 문자가 입력됨", "Ctrl C 후 한 줄씩 다시 입력"],
        ["curl URL에 Markdown 링크 포함", "대화 화면의 표시용 링크를 그대로 복사", "http부터 시작하는 순수 URL만 입력"],
        ["find 괄호 syntax error", "괄호 escape 누락", "백슬래시 괄호를 사용"],
        ["sudo ssudo", "두 명령이 합쳐짐", "현재 process와 port를 확인하고 명령을 한 번만 실행"],
    ], widths=[4.3, 5.2, 6.0])
    b.h1("로그 수집 최소 세트")
    b.code("hostname; hostname -i; date '+%F %T %Z'\nps -ef | grep '<PROCESS_PATTERN>'\nsudo ss -lntp | grep '<PORT>'\ngetent hosts '<HOSTNAME>'\ntail -n 200 '<SERVICE_LOG>'\ngrep -nE '<SETTING>' '<CONFIG_FILE>'")
    b.p("장애를 전달할 때는 실행한 명령, 전체 오류 문구, 발생 시각, 서버 hostname과 IP, 직전 변경 내용, 정상 서버의 비교 결과를 함께 남긴다.")
    return b.save()


paths = [build_runbook(), build_checklist(), build_troubleshooting()]
for path in paths:
    print(path)
