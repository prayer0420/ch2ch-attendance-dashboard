from pathlib import Path
from datetime import date
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "산출물"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = OUT_DIR / "메리츠_xen02_xen01_1차_이관실습_전체기록.docx"


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)
sec.header_distance = Cm(0.8)
sec.footer_distance = Cm(0.8)


def set_font(run, name="맑은 고딕", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def style_font(style, name="맑은 고딕", size=10.2, bold=False, color=(0, 0, 0)):
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(*color)


styles = doc.styles
style_font(styles["Normal"], size=10.2)
styles["Normal"].paragraph_format.line_spacing = 1.18
styles["Normal"].paragraph_format.space_after = Pt(5)

style_font(styles["Title"], size=25, bold=True)
styles["Title"].paragraph_format.space_after = Pt(14)
styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

for name, size, before, after in [
    ("Heading 1", 17, 18, 8),
    ("Heading 2", 13.5, 13, 6),
    ("Heading 3", 11.5, 9, 4),
]:
    style_font(styles[name], size=size, bold=True)
    styles[name].paragraph_format.space_before = Pt(before)
    styles[name].paragraph_format.space_after = Pt(after)
    styles[name].paragraph_format.keep_with_next = True

if "Code Block" not in styles:
    code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
else:
    code_style = styles["Code Block"]
style_font(code_style, name="Consolas", size=8.2, color=(25, 25, 25))
code_style.paragraph_format.left_indent = Cm(0.35)
code_style.paragraph_format.right_indent = Cm(0.25)
code_style.paragraph_format.space_before = Pt(3)
code_style.paragraph_format.space_after = Pt(7)
code_style.paragraph_format.line_spacing = 1.0

if "Small Note" not in styles:
    note_style = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
else:
    note_style = styles["Small Note"]
style_font(note_style, size=9, color=(70, 70, 70))
note_style.paragraph_format.left_indent = Cm(0.35)
note_style.paragraph_format.space_after = Pt(5)


def shade_paragraph(paragraph, fill="F3F5F7"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_title(text):
    p = doc.add_paragraph(style="Title")
    p.add_run(text)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)
    return p


def add_h1(text, page_break=False):
    if page_break:
        doc.add_page_break()
    text = re.sub(r"[^0-9A-Za-z가-힣 ]+", " ", text)
    return doc.add_heading(re.sub(r"\s+", " ", text).strip(), level=1)


def add_h2(text):
    text = re.sub(r"[^0-9A-Za-z가-힣 ]+", " ", text)
    return doc.add_heading(re.sub(r"\s+", " ", text).strip(), level=2)


def add_h3(text):
    text = re.sub(r"[^0-9A-Za-z가-힣 ]+", " ", text)
    return doc.add_heading(re.sub(r"\s+", " ", text).strip(), level=3)


def add_p(text="", bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        r1.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_note(text):
    p = doc.add_paragraph(style="Small Note")
    p.add_run(text)
    return p


def add_bullets(items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def add_numbers(items):
    for idx, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.add_run(f"{idx}. {item}")


def add_code(text):
    p = doc.add_paragraph(style="Code Block")
    shade_paragraph(p)
    lines = text.strip("\n").splitlines() or [""]
    for i, line in enumerate(lines):
        r = p.add_run(line)
        set_font(r, name="Consolas", size=8.2)
        if i < len(lines) - 1:
            r.add_break()
    return p


def add_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for idx, h in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h))
        set_font(r, size=9, bold=True, color=(255, 255, 255))
        if widths:
            cell.width = Cm(widths[idx])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            if ridx % 2 == 1:
                set_cell_shading(cell, "F6F9FC")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_font(r, size=8.8)
            if widths:
                cell.width = Cm(widths[idx])
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_step(title, purpose, command=None, result=None, issue=None, resolution=None):
    add_h3(title)
    add_p("목적: " + purpose, bold_lead="목적:")
    if command:
        add_code(command)
    if result:
        add_p("관찰 결과: " + result, bold_lead="관찰 결과:")
    if issue:
        add_p("문제와 원인: " + issue, bold_lead="문제와 원인:")
    if resolution:
        add_p("조치: " + resolution, bold_lead="조치:")


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


# Footer
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = fp.add_run("메리츠 xen02 to xen01 1차 이관 실습 전체 기록   ")
set_font(fr, size=8, color=(95, 95, 95))
add_page_number(fp)


# Cover
add_title("메리츠 xen02 to xen01 1차 이관 실습 전체 기록")
p = doc.add_paragraph()
r = p.add_run("MariaDB MaxScale Redis ZooKeeper Solr Elasticsearch 애플리케이션 HAProxy")
set_font(r, size=13, bold=True, color=(31, 78, 121))
p.paragraph_format.space_after = Pt(18)

add_p("이 문서는 2026년 9월 10일 수행한 3노드 테스트 환경 이관 실습의 전체 기록이다. 단순 성공 절차만 남기지 않고, 실제로 발생한 인증 오류, 복제 중단, 호스트 해석 문제, 권한 오류, 포트 충돌, 클러스터 상태 불일치와 잘못된 명령 입력까지 함께 정리했다. 다음 실습에서는 각 단계의 목적과 정상 기준을 이해하고, 같은 오류가 발생했을 때 원인을 빠르게 좁히는 데 사용한다.")
add_p("이번 실습은 xen02의 192.168.1.101부터 103까지 구성된 환경을 xen01의 192.168.1.81부터 83까지 옮겨 독립적으로 기동한 작업이다. 최종적으로 MariaDB와 MaxScale, Redis 6노드 클러스터, ZooKeeper 3노드 ensemble, SolrCloud, 메리츠용 Elasticsearch 7.8.0, HAProxy와 주요 애플리케이션을 대상 환경에서 확인했다.")

add_table(
    ["구분", "내용"],
    [
        ["실습 일자", "2026-09-10"],
        ["원본 환경", "xen02.01 to xen02.03, 192.168.1.101 to 192.168.1.103"],
        ["대상 환경", "xen01.01 to xen01.03, 192.168.1.81 to 192.168.1.83"],
        ["문서 성격", "전체본 1차 초안, 시행착오와 후속 점검 포함"],
        ["보안 처리", "비밀번호와 인증키는 실제 값 대신 자리표시자로 표기"],
    ],
    widths=[3.3, 13.7],
)

add_h1("문서 사용 방법")
add_bullets([
    "각 장의 앞부분에서 구성과 목적을 먼저 이해한 뒤 명령을 실행한다.",
    "명령 결과는 정상 기준과 비교한다. 포트가 열렸다는 사실만으로 클러스터 전체가 정상이라고 판단하지 않는다.",
    "source와 target을 항상 hostname과 IP로 재확인한다. 삭제, 초기화, FLUSH, RESET, 복제 변경 명령은 대상 서버임을 확인한 뒤 실행한다.",
    "문서의 비밀번호 자리표시자는 현장 비밀관리 기준에 따라 실제 값을 주입한다. 쉘 기록과 문서에 실제 비밀번호를 남기지 않는다.",
])

add_h2("목차")
toc_items = [
    "1 이관 범위와 최종 구조",
    "2 공통 준비와 접속 문제",
    "3 MariaDB와 MaxScale",
    "4 Redis Cluster",
    "5 ZooKeeper",
    "6 SolrCloud",
    "7 공유 NFS 백업 경로",
    "8 Elasticsearch 7.8.0",
    "9 Logstash 확인",
    "10 애플리케이션 이관과 설정 변경",
    "11 HAProxy와 웹 접속",
    "12 학습배포 실패 분석",
    "13 최종 종합 검증",
    "14 미완료 항목과 다음 실습 개선사항",
    "15 오류별 트러블슈팅 표",
    "16 명령어 입력 실수와 해석 방법",
]
add_numbers(toc_items)


add_h1("1 이관 범위와 최종 구조", page_break=True)
add_h2("1.1 서버 대응 관계")
add_table(
    ["역할", "원본 xen02", "대상 xen01", "주요 별칭"],
    [
        ["1번 노드", "192.168.1.101", "192.168.1.81", "mariadb1 redis1 zk1 solr1 es1"],
        ["2번 노드", "192.168.1.102", "192.168.1.82", "mariadb2 redis2 zk2 solr2 es2"],
        ["3번 노드", "192.168.1.103", "192.168.1.83", "mariadb3 redis3 zk3 solr3 es3"],
    ],
    widths=[2.6, 3.6, 3.6, 7.2],
)

add_h2("1.2 서비스와 포트")
add_table(
    ["서비스", "구성", "대상 포트", "최종 확인"],
    [
        ["MaxScale", "xen01.01 listener", "3306", "접속 성공"],
        ["MariaDB", "1 master 2 slave", "13306", "MaxScale에서 Running"],
        ["Redis", "3 master 3 replica", "7000 7001 17000 17001", "16384 slots 정상"],
        ["ZooKeeper", "1 leader 2 follower", "2181 2888 3888", "3대 imok"],
        ["SolrCloud", "3 nodes", "8983", "6 collections, 22 active states"],
        ["Meritz Elasticsearch", "3 nodes, 7.8.0", "9210 9310", "green"],
        ["기존 Elasticsearch", "xen01.02 중심 9.1.1", "9200 9300", "보존"],
        ["HAProxy", "xen01.01", "80", "cms login 200"],
        ["Applications", "5개 공통, scheduler 1개", "8081 8180 8280 8380 8480 8580", "xen01.01 확인"],
    ],
    widths=[3.0, 4.0, 4.0, 6.0],
)

add_h2("1.3 이번 실습의 전환 방식")
add_bullets([
    "MariaDB는 복제 구조를 정상화하고 MaxScale이 대상 DB에 접속하도록 구성했다.",
    "Redis는 runtime을 복사한 뒤 빈 3 master 3 replica 클러스터를 구성했다. 일부 cache key는 시험용으로 옮겼으며, 최종적으로 애플리케이션 실행 과정에서 key가 증가했다.",
    "ZooKeeper는 3노드 ensemble을 대상 주소로 구성하고 Solr 전용 chroot를 분리했다.",
    "Solr는 configset과 모든 collection을 API BACKUP과 RESTORE 방식으로 이관했다.",
    "Elasticsearch는 원본 7.8.0을 대상의 별도 포트 9210과 9310에 구성하고 Snapshot과 Restore로 데이터를 옮겼다.",
    "기존 xen01.02의 Elasticsearch 9.1.1은 9200과 9300에서 그대로 유지했다.",
    "애플리케이션은 기존 JAR checksum을 유지하고 대상 의존 서비스 주소와 포트를 조정했다.",
])


add_h1("2 공통 준비와 접속 문제", page_break=True)
add_h2("2.1 source와 target 구분")
add_code("""hostname
hostname -i
cat /etc/os-release
date '+%F %T %Z'""")
add_p("같은 admin 프롬프트를 사용하므로 명령 실행 전 hostname과 IP를 함께 확인해야 한다. 실제 실습 중 xen01 프롬프트만 보고 작업했지만 hostname -i가 192.168.1.82로 나온 사례가 있었고, 이 때문에 master에서 실행할 명령이 replica에서 실행되었다.")

add_h2("2.2 hosts 별칭")
add_code("""getent hosts mariadb mariadb1 mariadb2 mariadb3
getent hosts redis1 redis2 redis3
getent hosts zk1 zk2 zk3
getent hosts solr1 solr2 solr3
getent hosts es1 es2 es3
getent hosts master engine1 engine2 engine3 scheduler1""")
add_p("출력이 없으면 DNS 또는 /etc/hosts에 별칭이 없다. 이번 실습에서는 Redis, ZooKeeper, Solr, Elasticsearch와 애플리케이션 이름을 대상 IP로 등록했다. source와 target이 동시에 켜져 있으므로 각 환경의 /etc/hosts가 자기 환경의 IP를 가리켜야 한다.")
add_code("""192.168.1.81 redis1 zk1 solr1 es1 mariadb1 engine1 gateway1 cms1 master1 chat-ui1 scheduler1
192.168.1.82 redis2 zk2 solr2 es2 mariadb2 engine2 gateway2 cms2 master2 chat-ui2
192.168.1.83 redis3 zk3 solr3 es3 mariadb3 engine3 gateway3 cms3 master3 chat-ui3""")
add_note("실제 /etc/hosts는 기존 항목과 중복되지 않도록 편집한다. 한 이름이 source와 target 두 IP에 동시에 등록되면 결과가 일정하지 않을 수 있다.")

add_h2("2.3 SSH와 파일 전송")
add_p("서버 간 scp는 비밀번호 인증이 막혀 있고 송신 서버에 대상 인증용 키가 없어 Permission denied publickey로 실패했다. 운영 정책상 서버 간 신규 SSH 키를 만들지 않고, 수신 측에서 가능한 경로로 가져오거나 터미널 프로그램의 파일 전송 기능 또는 공유 NFS를 사용했다.")
add_code("""scp /tmp/파일.tar.gz admin@192.168.1.82:/tmp/
# 결과
Permission denied (publickey,gssapi-keyex,gssapi-with-mic).""")
add_p("이 메시지는 파일이나 경로 문제보다 SSH 인증 정책 문제를 의미한다. 파일 크기와 checksum은 전송 후 반드시 다시 확인했다.")
add_code("""ls -lh /tmp/파일.tar.gz
sha256sum /tmp/파일.tar.gz""")

add_h2("2.4 디렉터리 권한의 반복 문제")
add_p("runtime을 root 권한으로 압축 해제하거나 새 디렉터리를 root로 만들면 admin으로 실행되는 서비스가 data, log, pid 디렉터리에 쓰지 못했다. Redis, ZooKeeper, Solr, Elasticsearch와 애플리케이션마다 동일한 유형의 문제가 반복됐다.")
add_code("""sudo install -d -o admin -g admin -m 750 /logs/서비스명
sudo chown -R admin:admin /application/제품/data /application/제품/logdir
sudo -u admin test -w /application/제품/data && echo WRITABLE""")

add_h2("2.5 네트워크와 방화벽")
add_p("No route to host는 서비스 프로세스보다 방화벽 또는 네트워크 경로를 먼저 의심해야 하고, Connection refused는 대상 호스트까지 도달했으나 해당 포트가 LISTEN하지 않는 상태를 의미한다.")
add_code("""nc -zvw2 대상호스트 포트
sudo ss -lntp | grep ':포트\\b'
sudo firewall-cmd --zone=public --list-all""")


add_h1("3 MariaDB와 MaxScale", page_break=True)
add_h2("3.1 초기 상태와 오류")
add_p("MaxScale 서비스는 시작됐지만 mxs_service 계정으로 backend의 사용자 정보를 조회하지 못했다.")
add_code("""Error 1045: Access denied for user 'mxs_service'@'192.168.1.81'
Error 1130: Host '192.168.1.81' is not allowed to connect""")
add_p("1045는 계정 또는 비밀번호가 맞지 않는 경우이고, 1130은 접속한 client host에 대응하는 계정 행이 없거나 host 조건이 맞지 않는 경우다. MariaDB 계정은 사용자 이름과 Host의 조합이므로 mxs_service@192.168.1.81을 정확히 다뤄야 한다.")

add_h2("3.2 계정 확인")
add_code('''/mariadb/bin/mariadb -S /mariadb/run/mariadb.sock -e \\
"SELECT @@server_id; SELECT User,Host,plugin FROM mysql.user WHERE User='mxs_service';"''')
add_p("이 단계에서 hostname -i와 @@server_id를 함께 확인했다. server_id 202가 나온 노드는 replica 2번이므로 master로 착각하지 않도록 했다.")

add_h2("3.3 ALTER USER와 복제 SQL thread 중단")
add_code("""ERROR 1396: Operation ALTER USER failed for 'mxs_service'@'192.168.1.81'
ERROR 1133: Can't find any matching row in the user table""")
add_p("mysql.user 조회에는 행이 보였지만 SHOW CREATE USER와 SET PASSWORD는 행을 찾지 못했다. 이어서 SHOW SLAVE STATUS에서 SQL thread가 ALTER USER 이벤트에서 멈춘 사실을 확인했다.")
add_code("""/mariadb/bin/mariadb -S /mariadb/run/mariadb.sock -e "SHOW SLAVE STATUS\\G" | \\
grep -E 'Slave_IO_Running|Slave_SQL_Running|Last_SQL_Error|Seconds_Behind_Master|Gtid_IO_Pos'""")
add_code("""Slave_IO_Running: Yes
Slave_SQL_Running: No
Last_SQL_Error: Operation ALTER USER failed for 'mxs_service'@'192.168.1.81'""")

add_h2("3.4 sql_log_bin 변수 사용 오류")
add_code("""SET GLOBAL sql_log_bin=OFF;
ERROR 1238: Variable 'sql_log_bin' is a SESSION variable""")
add_p("sql_log_bin은 현재 세션에서만 변경할 수 있으므로 동일한 MariaDB client 세션 안에서 SET SESSION을 사용했다.")
add_code("""SET SESSION sql_log_bin=OFF;
CREATE OR REPLACE USER 'mxs_service'@'192.168.1.81' IDENTIFIED BY '<MXS_SERVICE_PASSWORD>';
GRANT SHOW DATABASES ON *.* TO 'mxs_service'@'192.168.1.81';
GRANT SELECT ON mysql.* TO 'mxs_service'@'192.168.1.81';
SET SESSION sql_log_bin=ON;""")
add_note("비밀번호는 문서에 저장하지 않는다. CREATE OR REPLACE USER는 계정을 바꾸므로 source와 target, 복제 방향을 확인한 뒤 실행한다.")

add_h2("3.5 GTID strict mode 오류")
add_code("""An attempt was made to binlog GTID 1-201-6061 which would create an out-of-order
sequence number with existing GTID 1-202-6061, and gtid strict mode is enabled""")
add_p("replica에서 로컬 계정 DDL이 GTID domain과 sequence에 영향을 주면서 source GTID 적용 순서와 충돌했다. 이 오류는 단순 START SLAVE 반복으로 해결할 문제가 아니다. 로컬 관리 계정 변경은 sql_log_bin=OFF 세션에서 수행하고, GTID 상태와 relay 적용 위치를 확인한 뒤 SQL thread를 재개해야 한다.")

add_h2("3.6 repl_user 오류")
add_p("mariadb3에서는 repl_user@192.168.1.% ALTER USER 이벤트에서도 SQL thread가 멈췄다. 계정 DDL이 복제될 때 각 노드의 기존 사용자 상태가 다르면 같은 문제가 재발할 수 있음을 확인했다.")
add_code("""Last_SQL_Error: Operation ALTER USER failed for 'repl_user'@'192.168.1.%'""")

add_h2("3.7 MaxScale 런타임 계정 변경")
add_code("""sudo maxctrl alter service MariaDB-Service \\
user=mxs_service password='<MXS_SERVICE_PASSWORD>'""")
add_p("이 명령은 runtime 변경을 /mariadb/maxscale/maxscale.cnf.d/MariaDB-Service.cnf에 저장하며 기존 정적 설정을 override한다. 재기동 후에도 유지되므로 별도 파일이 생성됐다는 경고는 실패가 아니다.")

add_h2("3.8 최종 검증")
add_code("""maxctrl list servers
nc -zvw3 mariadb 3306""")
add_code("""mariadb1  Master, Running  GTID 1-201-6135
mariadb2  Slave, Running   GTID 1-201-6135
mariadb3  Slave, Running   GTID 1-201-6067""")
add_p("최종 종합 검증 시점에 mariadb3의 GTID가 master와 다르게 보였다. MaxScale은 Running으로 표시했지만 완전 동기화 판정은 SHOW SLAVE STATUS에서 IO와 SQL thread가 Yes이고 Seconds_Behind_Master가 0인지 확인한 뒤 내려야 한다. 이 항목은 후속 확인 대상으로 남긴다.")


add_h1("4 Redis Cluster", page_break=True)
add_h2("4.1 구성 이해")
add_p("각 VM에 Redis 프로세스가 두 개씩 존재한다. 7000은 master 후보, 7001은 replica 후보이며 3대 합계 6노드다. 한 VM에 Redis가 하나씩만 있다는 뜻이 아니라, 각 VM에 master와 replica 역할의 프로세스가 함께 존재한다.")
add_table(
    ["VM", "7000", "7001", "cluster bus"],
    [
        ["redis1 192.168.1.81", "master slots 0 to 5460", "redis3 master의 replica", "17000 17001"],
        ["redis2 192.168.1.82", "master slots 5461 to 10922", "redis1 master의 replica", "17000 17001"],
        ["redis3 192.168.1.83", "master slots 10923 to 16383", "redis2 master의 replica", "17000 17001"],
    ],
    widths=[4.0, 4.8, 5.2, 3.0],
)

add_h2("4.2 runtime 압축과 배포")
add_code("""sudo tar -czpf /tmp/redis_runtime.tar.gz \\
-C /application --exclude=redis/db redis
ls -lh /tmp/redis_runtime.tar.gz""")
add_p("실습에서 생성된 runtime archive는 약 50 MB였다. db 디렉터리는 제외해 기존 cluster state, AOF, RDB를 섞지 않았다. xen01.02에는 처음 /application 자체가 없어 runtime을 먼저 배포했다.")

add_h2("4.3 노드별 announce 주소")
add_code("""grep -E '^(port|dir|cluster-enabled|cluster-config-file|cluster-announce-ip|cluster-announce-port|cluster-announce-bus-port)' \\
/application/redis/conf/7000.conf /application/redis/conf/7001.conf""")
add_p("archive를 복사한 직후 모든 서버의 cluster-announce-ip가 redis1로 남아 있었다. 각 서버에서 redis1, redis2, redis3으로 수정했다. announce 주소가 잘못되면 다른 노드가 접속할 주소를 잘못 학습한다.")

add_h2("4.4 로그 디렉터리 권한")
add_code("""mkdir: cannot create directory '/logs/redis': Permission denied
redis 7001 stopped
redis 7000 stopped""")
add_code("""sudo install -d -o admin -g admin -m 750 /logs/redis
/application/redis/start.sh
/application/redis/status.sh""")
add_p("로그 디렉터리를 admin이 쓸 수 있게 만든 뒤 6개 프로세스가 모두 ready=yes로 기동됐다.")

add_h2("4.5 포트 상호 통신")
add_code("""for h in redis1 redis2 redis3; do
  for p in 7000 7001 17000 17001; do
    nc -zvw1 "$h" "$p"
  done
done""")
add_p("처음에는 redis2와 redis3에 No route to host와 TIMEOUT이 발생했다. 방화벽과 서비스 기동을 조정한 뒤 12개 조합이 모두 Connected로 확인됐다.")

add_h2("4.6 첫 cluster create 실패")
add_code("""[ERR] Node redis1:7000 is not empty.
Either the node already knows other nodes or contains some key in database 0.""")
add_p("기존 자동 cluster state나 key가 남아 있어 빈 노드 조건을 만족하지 못했다. target 노드임을 확인한 뒤 dynamic state와 data를 정리하고 빈 인스턴스로 재기동해야 한다. source에는 RESET이나 FLUSH를 실행하지 않는다.")

add_h2("4.7 hostname 주소 오류")
add_code("""ERR Invalid node address specified: redis1:7000""")
add_p("Redis 5 계열의 CLUSTER MEET 처리 과정에서 hostname 주소가 거부됐다. 최종 구성은 대상 IP를 사용해 생성했고, announce 결과도 IP로 확인했다.")
add_code("""/application/redis/bin/redis-cli --cluster create \\
192.168.1.81:7000 192.168.1.82:7000 192.168.1.83:7000 \\
192.168.1.81:7001 192.168.1.82:7001 192.168.1.83:7001 \\
--cluster-replicas 1""")

add_h2("4.8 key 이동 시행착오")
add_code("""for key in $(redis-cli -h 192.168.1.103 -p 7000 --scan); do
  redis-cli -h 192.168.1.103 -p 7000 \\
  MIGRATE 192.168.1.83 7000 "" 0 5000 COPY REPLACE KEYS "$key"
done""")
add_code("""NOKEY
(error) MOVED 6429 redis2:7000""")
add_p("cluster node에 직접 SCAN과 MIGRATE를 반복하면 key의 hash slot owner가 다른 경우 MOVED가 발생한다. TTL이 만료되면 NOKEY도 발생할 수 있다. 이후 target을 FLUSHDB ASYNC로 비우고 노드별 key를 다시 옮겼지만, Redis가 cache라는 런북 원칙상 운영 전환에서는 시스템 배포 기능으로 재생성하는 방식이 우선이다.")

add_h2("4.9 최종 상태")
add_code("""/application/redis/bin/redis-cli --cluster check redis1:7000
/application/redis/bin/redis-cli -h redis1 -p 7000 cluster info""")
add_code("""[OK] All nodes agree about slots configuration.
[OK] All 16384 slots covered.
cluster_state:ok
cluster_known_nodes:6
cluster_size:3""")
add_p("최종 종합 점검에서는 master별 key 수가 9, 10, 13으로 총 32개였다. 이전 이관 직후 26개보다 증가한 것은 애플리케이션과 cache 동작으로 key가 새로 만들어졌기 때문이다. cache 검증에서는 단순 count 일치보다 cluster 상태, TTL, 대표 기능과 오류 로그를 함께 확인한다.")


add_h1("5 ZooKeeper", page_break=True)
add_h2("5.1 source 실행 경로 찾기")
add_p("초기 find 명령으로 /application 아래 zoo.cfg가 나오지 않아 실행 프로세스에서 실제 경로를 역추적했다.")
add_code("""sudo ss -lntp | grep -E ':(2181|2888|3888)\\b'
ps -fp <ZOOKEEPER_PID>
sudo tr '\\0' ' ' < /proc/<ZOOKEEPER_PID>/cmdline; echo
sudo lsof -p <ZOOKEEPER_PID> | grep -Ei 'zoo|zookeeper|myid|cfg'""")
add_p("실제 프로세스는 Solr 패키지 안의 /application/solr/zk를 사용했고 QuorumPeerMain 설정은 /application/solr/zk/conf/zoo.cfg였다.")

add_h2("5.2 source 설정")
add_code("""tickTime=2000
initLimit=10
syncLimit=5
dataDir=/application/solr/zk/data
dataLogDir=/application/solr/zk/datalog
clientPort=2181
server.1=zk1:2888:3888
server.2=zk2:2888:3888
server.3=zk3:2888:3888""")
add_p("source 1번 노드의 myid는 1이었다. source 전체 확인 결과 follower 2대와 leader 1대, Node count 207이었다.")

add_h2("5.3 runtime 배포")
add_code("""sudo tar -czpf /tmp/zk_runtime.tar.gz \\
-C /application/solr --exclude=zk/data --exclude=zk/datalog zk""")
add_p("archive 크기는 약 24 MB였다. xen01.02와 xen01.03에는 /application/solr이 없었으므로 runtime을 먼저 전달하고 압축을 해제했다.")
add_code("""sudo mkdir -p /application/solr
sudo tar -xzpf /tmp/zk_runtime.tar.gz -C /application/solr
sudo mkdir -p /application/solr/zk/data /application/solr/zk/datalog
printf '2\\n' | sudo tee /application/solr/zk/data/myid""")
add_p("2번은 myid 2, 3번은 myid 3을 사용했다. 1번은 myid 1이었다.")

add_h2("5.4 로그 경로 오류")
add_code("""mkdir: cannot create directory '/logs/solr': Permission denied
FAILED TO START""")
add_code("""sudo install -d -o admin -g admin -m 750 /logs/solr/zk""")

add_h2("5.5 dataLogDir 권한 오류")
add_code("""Unable to create data directory /application/solr/zk/datalog/version-2
Unable to access datadir, exiting abnormally""")
add_code("""sudo chown -R admin:admin /application/solr/zk/data /application/solr/zk/datalog
sudo chmod -R u+rwX,go-rwx /application/solr/zk/data /application/solr/zk/datalog""")
add_p("zkServer.sh가 STARTED라고 출력해도 JVM이 곧 종료될 수 있다. status, ss, ruok와 로그를 함께 확인해야 한다.")

add_h2("5.6 nc 도구 부재")
add_code("""-bash: nc: command not found""")
add_p("nmap-ncat은 점검 도구일 뿐 ZooKeeper 실행 필수 패키지가 아니다. 설치하지 않고 Bash의 /dev/tcp로도 확인할 수 있다.")
add_code("""timeout 3 bash -c 'exec 3<>/dev/tcp/127.0.0.1/2181; printf "srvr" >&3; cat <&3' | \\
grep -E 'Mode|Node count'""")

add_h2("5.7 최종 상태")
add_code("""for h in zk1 zk2 zk3; do
  printf "ruok" | nc -w 3 "$h" 2181
  printf "srvr" | nc -w 3 "$h" 2181 | grep -E 'Mode|Node count'
done""")
add_code("""zk1 imok Mode: leader   Node count: 361
zk2 imok Mode: follower Node count: 361
zk3 imok Mode: follower Node count: 361""")
add_p("leader는 재선출에 따라 다른 노드로 바뀔 수 있다. 중요한 기준은 leader가 1대이고 follower가 2대이며 Node count가 동일한 것이다.")
add_p("실습에서는 xen01.01의 기존 version-2 snapshot이 남아 있는 상태가 관찰됐다. 이후 Solr를 별도 chroot /solr-xen01에 구성해 source 상태와 분리했지만, 다음 실습에서는 clean ensemble 원칙에 따라 세 노드 모두의 기존 dynamic data 보관과 초기화 여부를 명시적으로 기록해야 한다.")


add_h1("6 SolrCloud", page_break=True)
add_h2("6.1 runtime 존재 여부")
add_p("xen01.01에는 /application/solr/solr이 있었지만 xen01.02와 xen01.03에는 처음 존재하지 않았다. Solr runtime archive를 약 358 MB로 만들고 두 노드에 배포했다.")
add_code("""ls -l /application/solr/solr/bin/solr
grep -nE '^(SOLR_HOST|SOLR_PORT|ZK_HOST)' /application/solr/solr/bin/solr.in.sh""")

add_h2("6.2 hosts와 노드 이름")
add_code("""192.168.1.81 solr1
192.168.1.82 solr2
192.168.1.83 solr3""")
add_p("SOLR_HOST는 1번 solr1, 2번 solr2, 3번 solr3으로 고유해야 한다. archive 복사 과정에서 1번 설정이 solr2로 보이거나 세 노드가 solr-xen01이라는 동일한 이름을 사용한 시행착오가 있었다. 같은 SOLR_HOST를 사용하면 ZooKeeper에서 하나의 live node처럼 충돌하고 CoreContainer 초기화가 실패할 수 있다.")

add_h2("6.3 ZooKeeper chroot 분리")
add_code('''SOLR_HOST="solr1"
SOLR_PORT="8983"
ZK_HOST="zk1:2181,zk2:2181,zk3:2181/solr-xen01"''')
add_p("처음에는 /solr를 사용해 기존 상태와 섞였고 down replica가 보였다. source와 target이 동시에 실행되는 실습이므로 target은 /solr-xen01 chroot를 사용했다.")

add_h2("6.4 포트 사용 중과 HTTP 500")
add_code("""Port 8983 is already being used by another process
HTTP_CODE=500
CoreContainer is either not initialized or shutting down""")
add_p("포트 사용 중 메시지는 Solr가 이미 실행 중일 수 있다는 뜻이다. 무조건 다시 start하지 않고 ss와 기존 PID를 확인했다. HTTP 500은 process와 port는 살아 있지만 잘못된 ZK chroot 또는 중복 SOLR_HOST 때문에 CoreContainer가 초기화되지 않은 상태였다. solr1, solr2, solr3과 /solr-xen01을 맞춘 뒤 해결됐다.")

add_h2("6.5 기동")
add_code("""/application/solr/solr/bin/solr start -cloud -p 8983 \\
-z zk1:2181,zk2:2181,zk3:2181/solr-xen01""")
add_p("open file limit 1024와 entropy low 경고가 출력됐지만 이번 테스트 기동은 완료됐다. 운영에서는 ulimit 65000 이상과 entropy 상태를 별도로 조정해야 한다.")

add_h2("6.6 configset 이관")
add_p("source의 collection이 사용하는 chat-base-config를 ZooKeeper에서 내려받고 target chroot로 올렸다.")
add_code("""# source
/application/solr/solr/bin/solr zk downconfig \\
-n chat-base-config \\
-d /tmp/solr_config_export/chat-base-config \\
-z zk1:2181,zk2:2181,zk3:2181/solr

# target
/application/solr/solr/bin/solr zk upconfig \\
-n chat-base-config \\
-d /tmp/solr_config_import/chat-base-config \\
-z zk1:2181,zk2:2181,zk3:2181/solr-xen01""")
add_p("이 작업은 검색 문서가 아니라 schema.xml, solrconfig.xml, 한국어 분석 설정 등 collection 생성 규칙을 옮긴 것이다.")

add_h2("6.7 collection 백업과 복구")
add_p("source에는 6개 collection이 있었고 replicationFactor는 collection별로 2 또는 3이었다. 공유 NFS 경로에서 Solr Collections API BACKUP과 RESTORE를 수행했다.")
add_code("""# source backup 예시
curl -sS --max-time 300 \\
'http://127.0.0.1:8983/solr/admin/collections?action=BACKUP&collection=TC_CALLBOT_B&name=TC_CALLBOT_B_20260910&location=/data/solr_backup/TC_CALLBOT_B&wt=json'

# target restore 예시
curl -sS --max-time 300 \\
'http://127.0.0.1:8983/solr/admin/collections?action=RESTORE&collection=TC_CALLBOT_B&name=TC_CALLBOT_B_20260910&location=/data/solr_backup/TC_CALLBOT_B&wt=json'""")
add_p("TC_CALLBOT_B 복원 결과 replicationFactor 3의 replica가 solr1, solr2, solr3에 active로 생성됐고 문서 수 20개가 확인됐다.")

add_h2("6.8 복사된 core 디렉터리 문제")
add_p("xen01.02와 xen01.03의 Solr runtime을 통째로 복사했을 때 동일한 core.properties와 core 디렉터리가 양쪽에 존재했다. 이는 각 노드의 고유 replica 배치와 충돌할 수 있다. target chroot를 새로 구성하고 RESTORE가 생성한 replica만 사용하도록 정리한 뒤 상태를 재확인했다.")

add_h2("6.9 down replica와 통신")
add_code('''"state":"down"
"node_name":"solr2:8983_solr"
"node_name":"solr3:8983_solr"''')
add_p("처음에는 replica가 down으로 보였고 분산 query가 timeout됐다. 노드 이름, ZK chroot, core 디렉터리, 서버 간 8983 통신과 방화벽을 확인한 뒤 모든 replica가 active가 됐다.")

add_h2("6.10 최종 collection과 문서 수")
add_table(
    ["Collection", "source 문서 수", "target 문서 수", "결과"],
    [
        ["TC_CALLBOT_LOG", "0", "0", "일치"],
        ["TC_CALLBOT_B", "20", "20", "일치"],
        ["tc_01_A", "9", "9", "일치"],
        ["TC_CALLBOT_TEST", "0", "0", "일치"],
        ["tc_01_B", "1", "1", "일치"],
        ["TC_CALLBOT_A", "20", "20", "일치"],
    ],
    widths=[5.0, 3.4, 3.4, 3.5],
)
add_code('''curl -sS 'http://solr1:8983/solr/admin/collections?action=CLUSTERSTATUS&wt=json' | \\
grep -oE '"state":"[^"]+"' | sort | uniq -c

22 "state":"active"''')

add_h2("6.11 alias 확인")
add_code("""TC_CALLBOT -> TC_CALLBOT_A
tc_01      -> tc_01_A""")
add_p("source와 target의 alias가 동일함을 확인했다. collection backup과 별개로 alias 검증이 필요하다.")


add_h1("7 공유 NFS 백업 경로", page_break=True)
add_h2("7.1 구성 목적")
add_p("Solr와 Elasticsearch의 backup repository는 클러스터 모든 노드에서 동일한 절대 경로로 보여야 한다. source 192.168.1.101의 /data/solr_backup을 NFS로 export하고 source의 나머지 노드와 target 3대에 mount했다.")

add_h2("7.2 NFS server")
add_code("""sudo install -d -o admin -g admin -m 770 /data/solr_backup
sudo vi /etc/exports.d/solr-migration.exports
sudo systemctl enable --now nfs-server
sudo exportfs -rav
sudo exportfs -v""")
add_code("""/data/solr_backup 192.168.1.102(rw,sync,no_subtree_check,root_squash)
/data/solr_backup 192.168.1.103(rw,sync,no_subtree_check,root_squash)
/data/solr_backup 192.168.1.81(rw,sync,no_subtree_check,root_squash)
/data/solr_backup 192.168.1.82(rw,sync,no_subtree_check,root_squash)
/data/solr_backup 192.168.1.83(rw,sync,no_subtree_check,root_squash)""")

add_h2("7.3 방화벽과 mount 확인")
add_code("""sudo firewall-cmd --permanent --zone=public --add-service=nfs
sudo firewall-cmd --permanent --zone=public --add-service=mountd
sudo firewall-cmd --permanent --zone=public --add-service=rpc-bind
sudo firewall-cmd --reload

findmnt /data/solr_backup""")
add_p("target에서 SOURCE가 192.168.1.101:/data/solr_backup, FSTYPE이 nfs4, OPTIONS에 rw가 표시됐다.")

add_h2("7.4 쓰기 시험")
add_code("""touch /data/solr_backup/write-test-$(hostname -i | tr . -)
ls -l /data/solr_backup/write-test-*""")
add_p("101, 102, 103, 81, 82, 83의 test 파일이 source에서 모두 보였다. 이는 모든 노드가 같은 공유 경로를 보고 있다는 증거다.")

add_h2("7.5 root squash 권한 문제")
add_code("""sudo tar -xzpf /data/solr_backup/meritz_apps_common.tar.gz -C /application
tar: Cannot open: Permission denied""")
add_p("root_squash 때문에 client의 root가 NFS server에서 익명 사용자로 매핑됐다. archive와 상위 디렉터리는 admin 권한으로 읽을 수 있었지만 sudo root는 오히려 접근하지 못했다. 해결 방법은 admin으로 NFS에서 /tmp로 복사한 뒤 로컬 파일을 sudo로 압축 해제하는 것이다.")
add_code("""cp /data/solr_backup/meritz_apps_common.tar.gz /tmp/
sudo tar -xzpf /tmp/meritz_apps_common.tar.gz -C /application""")


add_h1("8 Elasticsearch 7.8.0", page_break=True)
add_h2("8.1 source 실제 버전 확인")
add_code("""curl -sS http://127.0.0.1:9200/ | grep -E '"name"|"cluster_name"|"number"'
curl -sS 'http://127.0.0.1:9200/_cat/nodes?v&h=name,ip,node.role,master,version'""")
add_p("xen02.01, 02, 03은 각각 es1, es2, es3이고 모두 7.8.0 trusted-context-es 클러스터였다. 문서에 별도 9.1.1 표기가 있었지만 실제 메리츠 서비스가 바라보는 9200은 7.8.0이었다.")

add_h2("8.2 target 기존 9.1.1 보존")
add_p("xen01.02에는 이미 Elasticsearch 9.1.1 tclab 클러스터가 9200과 9300을 사용하고 있었다. 이 서비스를 중지하거나 설정을 덮어쓰지 않고 메리츠 7.8.0을 별도 포트 9210과 9310으로 구성했다.")
add_table(
    ["대상", "버전", "클러스터", "HTTP", "Transport", "처리"],
    [
        ["기존 xen01.02", "9.1.1", "tclab", "9200", "9300", "유지"],
        ["이관 메리츠", "7.8.0", "trusted-context-es", "9210", "9310", "신규 3노드"],
    ],
    widths=[3.3, 2.1, 4.0, 2.0, 2.2, 2.5],
)

add_h2("8.3 runtime archive")
add_code("""sudo tar -czpf /tmp/elk78_runtime.tar.gz \\
-C /application --exclude=elk/data --exclude=elk/tmp elk
ls -lh /tmp/elk78_runtime.tar.gz
tar -tzf /tmp/elk78_runtime.tar.gz >/dev/null && echo ARCHIVE_OK""")
add_p("archive는 처음 확인할 때 1.2 GB에서 1.3 GB로 계속 증가해 작업이 멈춘 것처럼 보였다. tar가 아직 실행 중이었기 때문이며, 쉘 프롬프트가 돌아오고 ARCHIVE_OK를 확인한 뒤 전송했다.")

add_h2("8.4 target 설정")
add_code("""cluster.name: "trusted-context-es"
node.name: "es1"
path.data: "/application/elk-meritz/data"
path.logs: "/logs/elk-meritz"
network.host: 192.168.1.81
http.port: 9210
transport.port: 9310
discovery.seed_hosts: ["es1:9310", "es2:9310", "es3:9310"]
path.repo: ["/data/solr_backup"]""")
add_p("node.name과 network.host는 각 노드에 맞게 es1/81, es2/82, es3/83으로 설정했다. binary home은 /application/elk를 유지하고 working data와 log, pid만 별도 경로를 사용했다.")

add_h2("8.5 GC 로그 경로 오류")
add_code("""Error opening log file '/logs/elk/gc.log': No such file or directory
Invalid -Xlog option""")
add_p("elasticsearch.yml의 path.logs만 바꿔도 jvm.options 안의 HeapDumpPath, ErrorFile, GC log 경로는 자동으로 바뀌지 않는다. jvm.options의 /logs/elk를 /logs/elk-meritz로 변경했다.")
add_code("""sudo cp -a /application/elk/config/jvm.options \\
/application/elk/config/jvm.options.before-meritz
sudo sed -i 's#/logs/elk#/logs/elk-meritz#g' \\
/application/elk/config/jvm.options""")

add_h2("8.6 node lock 오류")
add_code("""failed to obtain node locks, tried [[/application/elk-meritz/data]]""")
add_p("data 경로가 쓰기 불가하거나 동일 경로를 사용하는 다른 Elasticsearch process 또는 stale lock이 있을 때 발생한다. PID와 process, data owner를 먼저 확인하고 같은 data 경로를 두 process가 사용하지 않도록 했다.")

add_h2("8.7 vm max map count")
add_code("""bootstrap checks failed
vm.max_map_count [65530] is too low, increase to at least [262144]""")
add_code("""sudo sysctl -w vm.max_map_count=262144
echo 'vm.max_map_count=262144' | sudo tee /etc/sysctl.d/99-elasticsearch.conf""")
add_p("운영에서는 일시적 sysctl 적용뿐 아니라 재부팅 후에도 유지되는 설정 파일을 확인한다.")

add_h2("8.8 localhost 접속 거부")
add_code('''curl http://127.0.0.1:9210/
Connection refused

curl http://192.168.1.81:9210/
"number" : "7.8.0"''')
add_p("network.host를 192.168.1.81로 제한했기 때문에 127.0.0.1에는 bind하지 않았다. ss에서 실제 bind 주소를 확인하고 노드 IP로 점검했다.")

add_h2("8.9 path repo 전체 노드 적용")
add_code("""location [/data/solr_backup/elasticsearch] doesn't match any of the locations
specified by path.repo because this setting is empty""")
add_p("처음에는 es1만 path.repo가 적용됐고 es2와 es3는 비어 있었다. filesystem repository는 모든 master/data 노드가 같은 path.repo를 가져야 한다. 세 노드 설정을 수정하고 순차 재기동한 뒤 Nodes Settings API에서 모두 확인했다.")

add_h2("8.10 source snapshot")
add_code("""curl -sS -X PUT -H 'Content-Type: application/json' \\
http://127.0.0.1:9200/_snapshot/meritz_migration_repo \\
-d '{"type":"fs","settings":{"location":"/data/solr_backup/elasticsearch"}}'

curl -sS -X PUT \\
'http://127.0.0.1:9200/_snapshot/meritz_migration_repo/meritz_20260910?wait_for_completion=true' \\
-H 'Content-Type: application/json' \\
-d '{"indices":"*","include_global_state":true}'""")
add_code("""state: SUCCESS
version: 7.8.0
shards total: 6
failed: 0
successful: 6""")
add_p("source cluster에 path.repo를 추가하기 위해 es1, es2, es3를 순차 재기동했다. snapshot에는 6개의 system index가 포함됐으며 전체 크기는 작았다.")

add_h2("8.11 target restore 결과")
add_code("""curl -sS 'http://es1:9210/_cluster/health?pretty'
curl -sS 'http://es1:9210/_cat/indices?expand_wildcards=all&v'""")
add_code("""status: green
number_of_nodes: 3
active_primary_shards: 6
active_shards: 12
unassigned_shards: 0""")
add_p("대상에는 .kibana_1, .kibana_task_manager_1, .kibana-event-log, APM system index와 ilm-history가 복원됐다.")

add_h2("8.12 repository verify 메서드")
add_code("""# 잘못된 GET
curl http://es1:9210/_snapshot/meritz_migration_repo/_verify?pretty
HTTP 405 allowed: POST

# 올바른 명령
curl -sS -X POST \\
'http://es1:9210/_snapshot/meritz_migration_repo/_verify?pretty'""")


add_h1("9 Logstash 확인", page_break=True)
add_p("xen02의 101, 102, 103에서 process, systemd service, 5044, 5000, 9600 포트와 설정 파일을 검색했지만 Logstash 실행 흔적을 찾지 못했다. xen02.03의 5000은 Docker Registry container의 포트였다.")
add_code("""ps -ef | grep '[l]ogstash'
sudo ss -lntp | grep -E ':(5044|5000|9600)\\b'
systemctl list-units --type=service --all | grep -Ei 'elastic|logstash'
sudo find /application /etc /opt -type f \\
\\( -name 'logstash.yml' -o -name 'pipelines.yml' -o -name '*.conf' \\) 2>/dev/null""")
add_p("이번 실습 범위에서는 실행 중인 메리츠 Logstash가 없다고 판단해 별도 이관을 진행하지 않았다. 운영 전환 전에는 애플리케이션 로그 전송 구조, Filebeat, container 또는 중앙 Logstash 사용 여부를 다시 확인한다.")


add_h1("10 애플리케이션 이관과 설정 변경", page_break=True)
add_h2("10.1 대상 배치 구조")
add_p("gateway, engine, master, cms, chat-ui는 3대에 공통 배치하는 구조이고 scheduler는 중복 실행을 막기 위해 1대에서만 운영하는 구조로 판단했다. 처음에는 xen01.01에만 애플리케이션 디렉터리가 있었고 xen01.02와 03에는 없었다.")

add_h2("10.2 archive 생성과 배포")
add_code("""tar -czpf /tmp/meritz_apps_common.tar.gz \\
-C /application gateway engine master cms chat-ui
ls -lh /tmp/meritz_apps_common.tar.gz""")
add_p("archive 크기는 약 1.9 GB였다. NFS에 직접 sudo tar로 쓰거나 읽을 때 root_squash로 Permission denied가 발생했으므로 admin으로 /tmp와 공유 경로 사이를 복사한 뒤 로컬에서 sudo 압축 해제했다.")

add_h2("10.3 JAR과 설정 무결성")
add_code("""sha256sum /application/cms/*.jar /application/master/*.jar
sha256sum /application/cms/conf/runtime.env /application/master/conf/runtime.env""")
add_p("source와 target의 CMS 및 Master JAR checksum이 동일했고 runtime.env checksum도 동일했다. Master application.yml 차이는 Elasticsearch port 9200에서 9210으로 바꾼 부분뿐이었다.")

add_h2("10.4 Elasticsearch endpoint 변경")
add_code("""# gateway
url: http://es1:9210
port: 9210

# master cms scheduler
elasticsearch:
  clustername: trusted-context-es
  host1: es1
  host2: es2
  host3: es3
  port: 9210
  scheme: http""")
add_p("application 파일의 첫 번째 server.port는 각 애플리케이션 자체 포트다. 예를 들어 gateway 8081, master 8380, cms 8280이다. Elasticsearch 하위의 port만 9210으로 변경했다. 백업 파일 .before-es9210에 남은 9200 문자열은 실행 설정이 아니므로 문제없다.")

add_h2("10.5 다른 의존 서비스")
add_p("애플리케이션은 Elasticsearch만 사용하는 것이 아니다. MariaDB/MaxScale, Redis, ZooKeeper/Solr, Master, Engine, Scheduler 이름을 모두 해석하고 접속할 수 있어야 한다.")
add_code("""nc -zvw2 mariadb 3306
redis-cli -h redis1 -p 7000 cluster info
printf 'ruok' | nc -w 3 zk1 2181
curl -sS http://solr1:8983/solr/admin/info/system?wt=json
curl -sS http://es1:9210/""")

add_h2("10.6 로그 디렉터리")
add_code("""sudo install -d -o admin -g admin -m 750 \\
/logs/master /logs/cms /logs/engine /logs/gateway /logs/chat-ui""")
add_p("xen01.02와 03에서 처음 start.sh를 실행했을 때 /logs 디렉터리를 만들 권한이 없어 모든 애플리케이션이 실패했다. 디렉터리를 admin 소유로 만든 뒤 재시도했다.")

add_h2("10.7 기동 순서")
add_numbers([
    "MariaDB와 MaxScale을 확인한다.",
    "Redis, ZooKeeper, SolrCloud, Elasticsearch를 확인한다.",
    "Master를 시작하고 8380 LISTEN을 확인한다.",
    "CMS를 시작하고 8280 LISTEN 및 /login을 확인한다.",
    "Engine, Gateway, Chat UI를 시작한다.",
    "Scheduler는 마지막에 승인된 1대에서만 시작한다.",
])
add_code("""/application/master/start.sh
/application/cms/start.sh
/application/engine/start.sh
/application/gateway/start.sh
/application/chat-ui/start.sh
/application/scheduler/start.sh""")

add_h2("10.8 Master 이름 해석과 CMS 기동 실패")
add_code("""Could not resolve host: master
Master app has not started yet
Cms application is stopping""")
add_p("CMS 설정의 master.api.url이 http://master였지만 target hosts에 master가 없어 CMS가 종료됐다. master 별칭을 대상 Master IP로 등록하고 Master가 완전히 올라온 뒤 CMS를 시작했다.")

add_h2("10.9 Master Auth key 403")
add_code("""HTTP_CODE=403
{"code":"9999","message":"Master Auth key error"}""")
add_p("인증 헤더 없이 Master 내부 API를 curl로 직접 호출해 403이 발생했다. source와 target JAR, runtime.env, config checksum을 비교했고, 이후 실제 CMS가 Master와 통신해 정상 기동했으므로 직접 curl의 403은 애플리케이션 장애 증거가 아니었다.")

add_h2("10.10 xen01.01 최종 프로세스")
add_code("""master    8380 running
cms       8280 running
engine    8180 running
gateway   8081 running
chat-ui   8480 running
scheduler 8580 running""")
add_p("status.sh 결과와 ss LISTEN을 함께 확인했다. status.sh가 PID만 찾는 방식일 수 있으므로 HTTP와 실제 업무 기능까지 확인해야 한다.")


add_h1("11 HAProxy와 웹 접속", page_break=True)
add_h2("11.1 source 구조 확인")
add_p("source xen02.01의 HAProxy는 80번 포트에서 Host 헤더에 따라 engine, gateway, cms, master, chat-ui backend로 분기했다. 각 backend는 101, 102, 103의 번호별 별칭과 애플리케이션 포트를 사용했다.")

add_h2("11.2 target 설치와 설정")
add_code("""sudo haproxy -c -f /etc/haproxy/haproxy.cfg
sudo systemctl enable --now haproxy
sudo systemctl restart haproxy
sudo ss -lntp | grep ':80\\b'""")
add_p("처음에는 haproxy 명령 자체가 없어 패키지를 설치했다. source 설정 파일을 target에 복사한 뒤 문법 검증에서 Configuration file is valid를 확인했고 0.0.0.0:80 LISTEN을 확인했다.")

add_h2("11.3 HTTP 503")
add_code("""HTTP_CODE=503
No server is available to handle this request.""")
add_p("HAProxy는 살아 있지만 backend CMS가 내려가 있거나 HAProxy가 cms1, cms2, cms3를 해석하거나 접속하지 못할 때 발생한다. backend 애플리케이션 포트와 hosts 등록을 확인한 뒤 해결했다.")

add_h2("11.4 Windows hosts")
add_code("""192.168.1.81 cms
192.168.1.81 meritz.easycms""")
add_p("Windows hosts에 기존 192.168.1.103 meritz.easycms와 신규 192.168.1.81 meritz.easycms가 동시에 존재해 중복 가능성이 있었다. 검증 시에는 하나의 이름이 하나의 target IP만 가리키도록 정리하고 ipconfig /flushdns를 실행했다.")

add_h2("11.5 ping 성공과 HTTP timeout")
add_p("ping cms가 192.168.1.81로 성공해도 TCP 80이 방화벽에서 막혀 있으면 웹은 timeout된다. target firewalld에서 HTTP 서비스를 허용한 뒤 웹 접속이 가능해졌다.")
add_code("""sudo firewall-cmd --permanent --zone=public --add-service=http
sudo firewall-cmd --reload
sudo firewall-cmd --zone=public --list-services""")

add_h2("11.6 최종 웹 확인")
add_code("""curl -sS -o /dev/null \\
-w 'HTTP_CODE=%{http_code}\\n' \\
-H 'Host: cms' \\
http://127.0.0.1/login

HTTP_CODE=200""")
add_p("root 경로에서 engine, gateway, master, chat-ui가 404를 반환하고 CMS가 401을 반환한 것은 각 애플리케이션의 유효 endpoint가 /가 아니기 때문일 수 있다. CMS의 실제 로그인 경로 /login에서 200을 확인했다.")


add_h1("12 학습배포 실패 분석", page_break=True)
add_h2("12.1 처음 의심한 항목")
add_p("웹 접속 후 학습배포가 실패해 Engine, Scheduler, CMS, Master 로그를 확인했다. Engine에는 js/common.js를 absolute file path로 해석할 수 없다는 오류가 있었지만 TC_CALLBOT과 tc_01의 Learning Engine 완료 로그와 EngineApplication 8180 기동 로그가 이어졌다. Scheduler의 RESERVE_DEPLOY_JOB과 LEARN_RECOMMEND_JOB 배치 로그 저장 결과도 true였다.")

add_h2("12.2 실제 실패 로그")
add_code("""alias name=B2B_BASIC_CODE, alias collection=null
null collection alias, result=null
Exception, Learning bot, botCode=B2B_BASIC_CODE
java.io.IOException: solr Data Learning or Delete Fail""")
add_p("CMS가 B2B_BASIC_CODE라는 Solr alias를 찾았지만 대응 collection이 없어 학습을 중단했다. source와 target의 LISTALIASES 결과는 모두 TC_CALLBOT과 tc_01 두 개뿐이었다.")

add_h2("12.3 source에서도 같은 실패")
add_code("""2026-09-03 source xen02
Cannot create channel messages, not exist channel, botCode=B2B_BASIC_CODE
Cannot create intent messages, not exist intent, botCode=B2B_BASIC_CODE
alias name=B2B_BASIC_CODE, alias collection=null
Exception, Learning bot, botCode=B2B_BASIC_CODE""")
add_p("동일 문제가 source에서도 이미 발생했으므로 이번 이관으로 생긴 장애가 아니다. B2B_BASIC_CODE에 임의로 TC_CALLBOT_A를 연결하면 다른 bot 데이터를 손상시킬 수 있으므로 alias를 추측해 만들지 않았다.")

add_h2("12.4 판단")
add_table(
    ["확인 항목", "결과", "판정"],
    [
        ["Solr 연결", "LISTALIASES 200 OK", "정상"],
        ["source alias", "TC_CALLBOT, tc_01", "target과 동일"],
        ["target alias", "TC_CALLBOT, tc_01", "source와 동일"],
        ["B2B_BASIC_CODE", "source와 target 모두 alias 없음", "기존 데이터 또는 업무 설정 문제"],
        ["이관 영향", "source에서도 동일 실패", "이관 장애 아님"],
    ],
    widths=[4.2, 7.4, 5.0],
)


add_h1("13 최종 종합 검증", page_break=True)
add_h2("13.1 포트와 프로세스")
add_code("""for ip in 81 82 83; do
  echo "===== xen01.$ip ====="
  ssh admin@192.168.1.$ip \\
  "hostname; hostname -i; sudo ss -lntp | grep -E ':(80|2181|2888|3888|7000|7001|8983|9210|9310)\\b' || true"
done""")
add_p("81, 82, 83에서 Redis, ZooKeeper, Solr, Meritz Elasticsearch 포트가 확인됐다. HAProxy 80은 xen01.01에서 확인됐다.")

add_h2("13.2 Redis")
add_code("""/application/redis/bin/redis-cli --cluster check redis1:7000
/application/redis/bin/redis-cli -h redis1 -p 7000 cluster info""")
add_p("정상 기준은 cluster_state ok, known nodes 6, master 3, replica 3, 전체 16384 slot covered, fail 또는 open slot 없음이다.")

add_h2("13.3 ZooKeeper")
add_code("""for h in zk1 zk2 zk3; do
  printf "ruok" | nc -w 3 "$h" 2181
  printf "srvr" | nc -w 3 "$h" 2181 | grep -E 'Mode|Node count'
done""")
add_p("실제 결과는 leader 1, follower 2, Node count 361 동일이었다.")

add_h2("13.4 Solr")
add_code("""curl -sS 'http://solr1:8983/solr/admin/collections?action=LIST&wt=json'
curl -sS 'http://solr1:8983/solr/admin/collections?action=CLUSTERSTATUS&wt=json' | \\
grep -oE '"state":"[^"]+"' | sort | uniq -c""")
add_p("6개 collection과 22 active 상태, source와 target 문서 수 일치를 확인했다.")

add_h2("13.5 MariaDB와 MaxScale")
add_code("""maxctrl list servers
nc -zvw3 mariadb 3306""")
add_p("listener 3306 접속과 master 1, slave 2 Running을 확인했다. mariadb3 GTID 차이는 별도 SHOW SLAVE STATUS 확인이 필요하다.")

add_h2("13.6 Elasticsearch")
add_code("""curl -sS 'http://es1:9210/_cat/nodes?v&h=name,ip,role,master,version'
curl -sS 'http://es1:9210/_cluster/health?pretty'
curl -sS http://192.168.1.82:9200/ | grep -E '"cluster_name"|"number"'""")
add_p("메리츠 7.8.0은 green 3노드이고, 기존 9.1.1 tclab도 9200에서 유지됐다.")

add_h2("13.7 애플리케이션과 웹")
add_code("""for d in master cms engine gateway chat-ui scheduler; do
  [ -x "/application/$d/status.sh" ] && /application/$d/status.sh
done
sudo ss -lntp | grep -E ':(8081|8180|8280|8380|8480|8580)\\b'
curl -sS -o /dev/null -w 'HTTP_CODE=%{http_code}\\n' \\
-H 'Host: cms' http://127.0.0.1/login""")
add_p("xen01.01에서 6개 애플리케이션 process와 port, CMS login HTTP 200을 확인했다.")

add_h2("13.8 최종 상태 요약")
add_table(
    ["영역", "최종 상태", "증거", "추가 확인"],
    [
        ["MariaDB MaxScale", "대체로 정상", "1 master 2 slave Running", "mariadb3 GTID와 lag"],
        ["Redis", "정상", "6 nodes, 16384 slots", "업무 cache 기능"],
        ["ZooKeeper", "정상", "leader 1 follower 2", "clean data 절차 보강"],
        ["SolrCloud", "정상", "6 collections, 22 active", "대표 한국어 검색"],
        ["Elasticsearch 7.8", "정상", "green, 3 nodes, 6 indices", "ILM SLM 정책 세부 비교"],
        ["Elasticsearch 9.1", "보존 정상", "tclab 9.1.1 응답", "운영 소유자 확인"],
        ["Applications xen01.01", "정상", "6 ports, CMS login 200", "업무 E2E"],
        ["Applications xen01.02 03", "부분 확인", "파일 배포와 일부 기동 수행", "노드별 status와 port"],
        ["학습배포 B2B", "실패", "source에서도 동일", "bot data와 alias 업무 확인"],
    ],
    widths=[3.5, 2.9, 6.0, 4.2],
)


add_h1("14 미완료 항목과 다음 실습 개선사항", page_break=True)
add_bullets([
    "mariadb3에서 SHOW SLAVE STATUS를 직접 확인하고 Slave IO Running, Slave SQL Running, Seconds Behind Master, Last SQL Error와 GTID를 증적으로 남긴다.",
    "xen01.02와 xen01.03에서 master, cms, engine, gateway, chat-ui의 status와 8081, 8180, 8280, 8380, 8480 LISTEN을 각각 확인한다.",
    "Scheduler가 한 노드에서만 실행되는지 81, 82, 83 전체 process 검색으로 확인한다.",
    "서비스가 수동 start.sh로만 기동된 상태라면 재부팅 후 자동 기동 체계를 별도 설계한다. PID 파일이 실제 process와 일치하는지도 확인한다.",
    "Solr open file limit 1024 경고를 운영 기준 65000 이상으로 조정한다.",
    "Elasticsearch vm.max_map_count 설정이 재부팅 후 유지되는지 확인한다.",
    "NFS는 migration 동안 단일 의존점이다. 최종 전환 후 유지 여부, readonly 전환, export 회수와 backup 보존 기간을 정한다.",
    "Windows hosts와 서버 hosts의 source IP 중복 항목을 정리하고 최종 DNS 또는 VIP 전환 절차를 확정한다.",
    "로그인, bot 조회, 학습, 배포, 대화 시뮬레이터, 검색, cache 재생성, 예약배포까지 업무 E2E 시험을 수행한다.",
    "B2B_BASIC_CODE는 source에서도 학습 실패하므로 별도 애플리케이션 데이터 결함으로 분리해 담당자에게 전달한다.",
    "최종 cutover에서는 source writer 중지와 마지막 MariaDB, Solr, Elasticsearch 동기화 시점을 기록한다. 이번 실습은 기능 검증용 1회 리허설이다.",
])


add_h1("15 오류별 트러블슈팅 표", page_break=True)
add_table(
    ["증상", "원인", "확인", "해결 방향"],
    [
        ["MariaDB 1045", "비밀번호 또는 계정 Host 불일치", "mysql.user, 접속 source IP", "정확한 User Host 계정과 MaxScale 비밀번호 일치"],
        ["MariaDB 1130", "client host 허용 계정 없음", "User Host 목록", "mxs_service@MaxScale IP 구성"],
        ["ALTER USER 1396", "replica의 계정 상태 불일치", "SHOW SLAVE STATUS", "복제 오류와 계정 DDL 순서 정리"],
        ["GTID out of order", "replica 로컬 DDL과 source GTID 충돌", "GTID IO Pos와 domain", "sql_log_bin OFF 세션과 GTID 복구 절차"],
        ["Redis node not empty", "key 또는 cluster state 잔존", "DBSIZE, CLUSTER NODES", "target data를 보관 후 빈 노드 구성"],
        ["Redis invalid node address", "hostname 주소 처리 문제", "redis-cli version", "대상 IP로 cluster create"],
        ["Redis MOVED", "key slot owner가 다른 node", "CLUSTER KEYSLOT", "cluster-aware 방식 또는 cache 재생성"],
        ["ZooKeeper STARTED 후 종료", "data 또는 datalog 쓰기 불가", "service log", "admin ownership과 write test"],
        ["nc command not found", "점검 도구 미설치", "command -v nc", "nmap-ncat 설치 또는 /dev/tcp"],
        ["Solr HTTP 500", "중복 SOLR_HOST 또는 잘못된 ZK state", "solr.in.sh, logs", "고유 node name과 target chroot"],
        ["Solr replica down", "노드 통신, stale core, ZK state", "CLUSTERSTATUS", "hosts firewall core 배치 점검"],
        ["NFS sudo Permission denied", "root squash", "exportfs -v, 파일 owner", "admin으로 로컬 복사 후 sudo 처리"],
        ["ES GC log 경로 오류", "jvm.options에 old path", "grep /logs/elk", "jvm.options와 log dir 수정"],
        ["ES node lock", "data 권한 또는 중복 process", "ps, lsof, data owner", "단일 process와 고유 data path"],
        ["ES bootstrap check", "vm.max_map_count 낮음", "sysctl", "262144 이상 적용"],
        ["ES localhost refused", "network.host가 node IP만 bind", "ss -lntp", "실제 node IP로 curl"],
        ["ES path.repo empty", "일부 node 미설정 또는 미재기동", "_nodes/settings", "3노드 설정 후 rolling restart"],
        ["HAProxy 503", "backend down 또는 이름 해석 실패", "HAProxy stats, nc", "backend process hosts firewall 확인"],
        ["CMS stopped", "Master 미기동 또는 master DNS 없음", "CMS log, getent", "Master 기동 확인 후 CMS 시작"],
        ["Master API 403", "인증 없는 직접 호출", "실제 CMS 통신", "403만으로 장애 판정하지 않음"],
        ["B2B 학습 실패", "Solr alias 없음", "LISTALIASES와 source log", "임의 alias 금지, 업무 데이터 확인"],
    ],
    widths=[3.5, 5.0, 3.5, 5.2],
)


add_h1("16 명령어 입력 실수와 해석 방법", page_break=True)
add_h2("16.1 출력 내용을 다시 명령으로 실행")
add_code("""22 "state":"active"
-bash: 22: command not found

"mode":"solrcloud"
-bash: mode:solrcloud: command not found""")
add_p("정상 결과 예시를 복사해 프롬프트에 붙이면 Bash는 그것을 명령으로 실행한다. command not found는 서비스 오류가 아니라 출력 예시를 실행한 입력 실수다.")

add_h2("16.2 프롬프트까지 복사")
add_code("""[admin@xen01 ~]$ 명령어
-bash: [admin@xen01: command not found""")
add_p("복사할 때 [admin@...]$ 부분을 제외하고 $ 뒤의 명령만 붙여 넣는다.")

add_h2("16.3 브래킷 붙여넣기 문자")
add_code("""^[[200~
-bash: $'\\E[200~': command not found""")
add_p("터미널의 bracketed paste 제어문자가 문자로 들어간 현상이다. Ctrl+C로 현재 줄을 취소하고 한 줄씩 다시 붙여 넣는다.")

add_h2("16.4 Markdown 링크가 섞인 URL")
add_p("대화 화면에서 http://주소가 [http://주소](http://주소) 형태로 복사되거나 &가 HTML escape로 보일 수 있다. 쉘에는 순수 URL만 입력한다.")
add_code("""curl -sS 'http://127.0.0.1:8983/solr/admin/collections?action=LIST&wt=json'""")

add_h2("16.5 잘못된 find 괄호")
add_code("""# 잘못된 예
find /application -type f ( -name '*.conf' )

# 올바른 예
find /application -type f \\( -name '*.conf' -o -name '*.yml' \\)""")

add_h2("16.6 sudo 오타와 반복 실행")
add_p("sudo ssudo처럼 명령이 합쳐지거나 이미 실행 중인 서비스를 다시 start한 사례가 있었다. 오류가 나오면 동일 명령을 반복하기보다 process, port, log를 먼저 확인한다.")
add_code("""ps -ef | grep '[프]로세스패턴'
sudo ss -lntp | grep ':포트\\b'
tail -n 100 /logs/서비스/서비스.log""")


add_h1("부록 A 다음 실습용 최소 시작 전 점검", page_break=True)
add_code("""hostname
hostname -i
date '+%F %T %Z'
df -hT / /data /application /tmp
getent hosts mariadb redis1 redis2 redis3 zk1 zk2 zk3 solr1 solr2 solr3 es1 es2 es3
sudo firewall-cmd --zone=public --list-all
sudo ss -lntp""")
add_p("이 결과를 source 3대와 target 3대에서 작업 시작 전에 보관하면 문서값과 현장값이 다른 문제를 일찍 발견할 수 있다.")

add_h1("부록 B 서비스 중지와 기동 순서")
add_h2("중지 순서")
add_numbers([
    "Scheduler와 모든 writer를 먼저 중지한다.",
    "Chat UI, Gateway, CMS, Engine, Master를 중지한다.",
    "Solr, Elasticsearch와 Redis를 중지한다.",
    "ZooKeeper를 중지한다.",
    "마지막으로 MaxScale과 MariaDB를 중지한다. 복제와 백업 절차에 맞춰 순서를 조정할 수 있다.",
])
add_h2("기동 순서")
add_numbers([
    "MariaDB master와 replica, MaxScale을 기동하고 복제를 확인한다.",
    "Redis Cluster를 기동하고 slot과 replica를 확인한다.",
    "ZooKeeper ensemble을 기동한다.",
    "SolrCloud를 기동하고 collection과 replica를 확인한다.",
    "Elasticsearch를 기동하고 green 상태를 확인한다.",
    "Master, CMS, Engine, Gateway, Chat UI를 순서대로 기동한다.",
    "업무 검증 후 Scheduler를 단 한 대에서 기동한다.",
])

add_h1("부록 C 보안과 기록 원칙", page_break=True)
add_bullets([
    "실제 비밀번호, SSH private key, database password, auth token을 문서와 화면 캡처에 남기지 않는다.",
    "변경 전 설정 파일은 timestamp 또는 목적을 포함한 이름으로 복사하고 checksum을 남긴다.",
    "source에서 파괴 명령을 실행하지 않는다. target의 FLUSH, RESET, data 이동도 hostname과 IP를 다시 확인한다.",
    "성공 판정은 process, port, cluster state, data count, HTTP와 실제 업무 기능의 복수 증거로 내린다.",
    "이번 문서는 1회 실습 기록이므로 다음 실습에서 실제 운영 절차로 축약하기 전에 미완료 항목과 승인 조건을 반영한다.",
])


# Core properties
doc.core_properties.title = "메리츠 xen02 to xen01 1차 이관 실습 전체 기록"
doc.core_properties.subject = "3노드 인프라 및 애플리케이션 이관과 트러블슈팅 기록"
doc.core_properties.author = "Migration Practice Team"
doc.core_properties.keywords = "MariaDB MaxScale Redis ZooKeeper Solr Elasticsearch HAProxy migration"

# Keep headings with following paragraphs and avoid orphan lines where possible.
for paragraph in doc.paragraphs:
    if paragraph.style.name.startswith("Heading"):
        paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.widow_control = True

doc.save(OUT_PATH)
print(OUT_PATH)
